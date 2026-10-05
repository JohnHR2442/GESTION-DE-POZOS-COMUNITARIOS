# -*- coding: utf-8 -*-
"""
Genera las tablas de operacionalizacion de variables con la estructura exacta
de la plantilla: encabezado "Definicion Operacional" combinado sobre
"Dimensiones" e "Indicadores".
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def style_cell(cell, text, bold=False, header=False, align='left'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
    }[align]
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.bold = bold
    if header:
        set_cell_bg(cell, 'D9D9D9')


def heading(doc, text, size=12):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    r.bold = True
    return p


def normal(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def build_table(doc, var_label, rows):
    """
    rows: lista de tuplas (variable, definicion_conceptual, dimension, indicador, escala)
    La estructura tiene 2 filas de encabezado:
      Fila 1: Variable | Definicion conceptual | Definicion Operacional (span 2) | Escala de Medicion
      Fila 2: (merge)  | (merge)               | Dimensiones | Indicadores     | (merge)
    """
    n_data = len(rows)
    total_rows = 2 + n_data
    table = doc.add_table(rows=total_rows, cols=5)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # --- Encabezados ---
    # Fila 0
    style_cell(table.cell(0, 0), var_label, bold=True, header=True, align='center')
    style_cell(table.cell(0, 1), 'Definición conceptual', bold=True, header=True, align='center')
    style_cell(table.cell(0, 2), 'Definición Operacional', bold=True, header=True, align='center')
    # celda (0,3) se combina con (0,2)
    style_cell(table.cell(0, 4), 'Escala de Medición', bold=True, header=True, align='center')

    # Fila 1 (subencabezados)
    style_cell(table.cell(1, 2), 'Dimensiones', bold=True, header=True, align='center')
    style_cell(table.cell(1, 3), 'Indicadores', bold=True, header=True, align='center')

    # Combinaciones verticales de encabezados (filas 0-1)
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))
    table.cell(0, 4).merge(table.cell(1, 4))
    # Combinacion horizontal "Definicion Operacional"
    table.cell(0, 2).merge(table.cell(0, 3))

    # --- Datos ---
    for i, (variable, defc, dim, ind, esc) in enumerate(rows):
        r = 2 + i
        style_cell(table.cell(r, 0), variable, align='left')
        style_cell(table.cell(r, 1), defc, align='left')
        style_cell(table.cell(r, 2), dim, align='left')
        style_cell(table.cell(r, 3), ind, align='left')
        style_cell(table.cell(r, 4), esc, align='center')

    # Combinar verticalmente la 1a columna (variable) y la 2a (definicion conceptual)
    first = table.cell(2, 0)
    for i in range(1, n_data):
        first = first.merge(table.cell(2 + i, 0))
    style_cell(table.cell(2, 0), rows[0][0], align='left')

    firstc = table.cell(2, 1)
    for i in range(1, n_data):
        firstc = firstc.merge(table.cell(2 + i, 1))
    style_cell(table.cell(2, 1), rows[0][1], align='left')

    return table


doc = Document()

# Estilo base del documento
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)

heading(doc, '1.5 Hipótesis y variables', 13)

normal(doc, 'En la presente investigación, de tipo descriptivo, transversal y '
            'correlacional, se puede suponer que:')

p = doc.add_paragraph()
r = p.add_run('Hipótesis de investigación (Hi): ')
r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(11)
r2 = p.add_run('La implementación de una aplicación móvil Android para la '
               'administración automática de turnos mejora significativamente el orden y la '
               'transparencia en la distribución del agua de los pozos comunitarios San Isidro, '
               'Zapata y Cárdenas.')
r2.font.name = 'Arial'; r2.font.size = Pt(11)
p.paragraph_format.line_spacing = 1.5
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p = doc.add_paragraph()
r = p.add_run('Hipótesis nula (H0): ')
r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(11)
r2 = p.add_run('La implementación de una aplicación móvil Android para la '
               'administración automática de turnos no influye significativamente en el orden y '
               'la transparencia en la distribución del agua de los pozos comunitarios San Isidro, '
               'Zapata y Cárdenas.')
r2.font.name = 'Arial'; r2.font.size = Pt(11)
p.paragraph_format.line_spacing = 1.5
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

heading(doc, 'Tabla de operacionalización de variables', 11)

# ---- Variable independiente ----
indep_rows = [
    ('Aplicación móvil para la administración de turnos de agua',
     'Sistema de software para dispositivos Android que automatiza el cálculo y control de los turnos, multas y avisos de los pozos comunitarios.',
     'Automatización del cálculo de turnos',
     'Porcentaje de turnos calculados automáticamente sin intervención manual',
     'Razón (%)'),
    ('', '', 'Disponibilidad del sistema',
     'Número de días con el sistema operativo y accesible', 'Razón (días)'),
    ('', '', 'Gestión de multas y avisos',
     'Número de multas y avisos registrados en el sistema', 'Razón (frecuencia)'),
    ('', '', 'Adopción por los usuarios',
     'Número de socios y contadores activos que utilizan la aplicación', 'Razón (frecuencia)'),
]
build_table(doc, 'Variable independiente', indep_rows)

doc.add_paragraph()

# ---- Variable dependiente ----
dep_rows = [
    ('Orden y transparencia en la distribución del agua',
     'Grado de claridad, equidad y ausencia de conflictos en la asignación y cumplimiento de los turnos de agua dentro de las comunidades.',
     'Reducción de conflictos',
     'Número de conflictos o reclamos por errores en los turnos',
     'Razón (frecuencia)'),
    ('', '', 'Claridad de la información',
     'Nivel de satisfacción de los usuarios con la información de turnos', 'Ordinal (1 a 5)'),
    ('', '', 'Exactitud de la rotación',
     'Porcentaje de turnos correctos respecto al orden acordado', 'Razón (%)'),
    ('', '', 'Transparencia en multas',
     'Porcentaje de multas con estado y motivo claramente registrado', 'Razón (%)'),
]
build_table(doc, 'Variable dependiente', dep_rows)

out = '/app/Tablas_Variables_Hipotesis.docx'
doc.save(out)
print('Guardado:', out)
