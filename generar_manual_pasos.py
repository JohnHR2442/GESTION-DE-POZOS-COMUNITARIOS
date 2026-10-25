from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CAP = ""
AZUL = "1F3A5F"
AZUL_CLARO = "DCE6F2"
fig_n = [0]
tab_n = [0]

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
pf = st.paragraph_format
pf.line_spacing = 1.5
pf.space_after = Pt(8)
for nombre, tam in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    h = doc.styles[nombre]
    h.font.name = "Arial"
    h.element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    h.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    h.font.size = Pt(tam)
    h.font.bold = True
    h.font.color.rgb = RGBColor.from_string(AZUL)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(8)
    h.paragraph_format.keep_with_next = True

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3), Cm(2.5)
sec.different_first_page_header_footer = True


def campo(par, instr, tam=10):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        r._r.append(e)
    r.font.size = Pt(tam)


def sombra(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def borde_inferior(par, color=AZUL):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "8"); bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), color)
    b.append(bot)
    pPr.append(b)


# Encabezado y pie
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = hp.add_run("Manual de Usuario – Aplicación móvil “Turnos de Pozo”")
r.font.size = Pt(9); r.bold = True
borde_inferior(hp)
fp = sec.footer.paragraphs[0]
fp.text = ""
from docx.enum.text import WD_TAB_ALIGNMENT
fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.09), WD_TAB_ALIGNMENT.RIGHT)
r = fp.add_run("Comisariado Ejidal de San José Carpinteros\t")
r.font.size = Pt(9)
campo(fp, "PAGE")


def p(texto, negrita_ini=None, centro=False, tam=None, cursiva=False, after=None):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.JUSTIFY
    if negrita_ini:
        rr = par.add_run(negrita_ini); rr.bold = True
        if tam: rr.font.size = Pt(tam)
    rr = par.add_run(texto)
    rr.italic = cursiva
    if tam: rr.font.size = Pt(tam)
    if after is not None: par.paragraph_format.space_after = Pt(after)
    return par


def vineta(texto, negrita_ini=None):
    par = doc.add_paragraph(style="List Bullet")
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(3)
    if negrita_ini:
        par.add_run(negrita_ini).bold = True
    par.add_run(texto)


def pasos(lista):
    for i, t in enumerate(lista, 1):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.left_indent = Cm(0.8)
        par.paragraph_format.first_line_indent = Cm(-0.8)
        par.paragraph_format.space_after = Pt(3)
        rr = par.add_run(f"Paso {i}.  "); rr.bold = True; rr.font.color.rgb = RGBColor.from_string(AZUL)
        par.add_run(t)


def nota(texto, titulo="Nota:", color="FFF4CE"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    sombra(c, color)
    par = c.paragraphs[0]
    par.paragraph_format.line_spacing = 1.15
    par.paragraph_format.space_after = Pt(2)
    rr = par.add_run(titulo + " "); rr.bold = True; rr.font.size = Pt(10)
    rr = par.add_run(texto); rr.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def figuras(items, alto=11.0):
    """items: lista de (archivo, titulo). Una o dos capturas por fila."""
    t = doc.add_table(rows=2, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (arch, titulo) in enumerate(items):
        fig_n[0] += 1
        c = t.cell(0, j)
        par = c.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(CAP + arch, height=Cm(alto))
        c2 = t.cell(1, j)
        cp = c2.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.line_spacing = 1.0
        cp.paragraph_format.space_after = Pt(6)
        cp.style = doc.styles["Caption"]
        rr = cp.add_run(f"Figura {fig_n[0]}. {titulo}")
        rr.font.size = Pt(9); rr.italic = False; rr.bold = False
        rr.font.name = "Arial"; rr.font.color.rgb = RGBColor(0, 0, 0)
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return fig_n[0] - len(items) + 1


def tabla(encabezados, filas, anchos, titulo):
    tab_n[0] += 1
    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cp.add_run(f"Tabla {tab_n[0]}. {titulo}")
    rr.font.size = Pt(9); rr.italic = False; rr.font.name = "Arial"; rr.font.color.rgb = RGBColor(0, 0, 0)
    cp.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, e in enumerate(encabezados):
        c = t.cell(0, j)
        sombra(c, AZUL)
        rr = c.paragraphs[0].add_run(e); rr.bold = True; rr.font.size = Pt(10); rr.font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, fila in enumerate(filas, 1):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            if i % 2 == 0:
                sombra(c, AZUL_CLARO)
            par = c.paragraphs[0]
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            rr = par.add_run(v); rr.font.size = Pt(10)
            if j == 0: rr.bold = True
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for fila in t.rows:
        for j, w in enumerate(anchos):
            fila.cells[j].width = Cm(w)
        trPr = fila._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
    t.autofit = False
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


H1 = lambda t: doc.add_heading(t, level=1)
H2 = lambda t: doc.add_heading(t, level=2)

# ======================= PORTADA =======================
for _ in range(3):
    doc.add_paragraph()
p("MANUAL DE USUARIO", centro=True, tam=26).runs[0].bold = True
p("Aplicación móvil “Turnos de Pozo”", centro=True, tam=16).runs[0].bold = True
p("Sistema Integral de Gestión de Pozos Comunitarios", centro=True, tam=13)
doc.add_paragraph()
p("Pozos comunitarios San Isidro, Zapata y Cárdenas", centro=True, tam=12)
p("Comisariado Ejidal de San José Carpinteros, Tepeaca, Puebla", centro=True, tam=12)
for _ in range(5):
    doc.add_paragraph()
t = doc.add_table(rows=4, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
datos = [("Elaboró:", "Juan José Huerta Rojas"), ("Versión del documento:", "1.0"),
         ("Plataforma:", "Android"), ("Fecha:", "Octubre de 2026")]
for i, (a, b) in enumerate(datos):
    ra = t.cell(i, 0).paragraphs[0].add_run(a); ra.bold = True
    t.cell(i, 1).paragraphs[0].add_run(b)
    t.cell(i, 0).width = Cm(5); t.cell(i, 1).width = Cm(7)
salto()

# ======================= ÍNDICE =======================
pi = doc.add_paragraph(); pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
ri = pi.add_run("CONTENIDO"); ri.bold = True; ri.font.size = Pt(14)
pt = doc.add_paragraph()
campo(pt, 'TOC \\o "1-2" \\h \\z \\u', tam=11)
pt.runs[0]._r.findall(qn("w:t"))[0].text = "Haga clic derecho aquí y seleccione “Actualizar campos” para mostrar el contenido."
salto()

# ======================= 1. INTRODUCCIÓN =======================
H1("1. Introducción")
p("El presente manual describe el uso de la aplicación móvil “Turnos de Pozo”, desarrollada para administrar "
  "de forma ordenada y transparente los turnos de agua de los pozos comunitarios San Isidro, Zapata y Cárdenas, "
  "pertenecientes al Ejido de San José Carpinteros. La aplicación sustituye el control manual en libretas y "
  "permite que cada socio consulte su turno, sus multas y los avisos del pozo desde su teléfono celular.")
H2("1.1 Objetivo del manual")
p("Guiar a los socios, contadores y usuarios en general, paso a paso, en el uso de cada una de las pantallas "
  "y funciones de la aplicación, de manera que puedan utilizarla sin necesidad de conocimientos técnicos.")
H2("1.2 Alcance")
p("Este documento abarca la instalación de la aplicación, el inicio de sesión, la recuperación de contraseña, "
  "las funciones del perfil Socio, las funciones del perfil Contador, la consulta pública de turnos y la "
  "solución de los problemas más comunes.")
H2("1.3 ¿A quién va dirigido?")
vineta(" ejidatarios con derecho de agua que necesitan conocer su turno de riego.", "Socios:")
vineta(" encargados de registrar multas y días sin servicio de cada pozo.", "Contadores:")
vineta(" cualquier persona de la comunidad que desee ver quién tiene el turno del día.", "Usuarios en general:")

# ======================= 2. REQUISITOS =======================
H1("2. Requisitos para usar la aplicación")
p("En la Tabla 1 se presentan los requisitos mínimos necesarios para instalar y utilizar la aplicación.")
tabla(["Requisito", "Descripción"],
      [("Dispositivo", "Teléfono celular o tableta con sistema operativo Android 7.0 o superior."),
       ("Conexión", "Acceso a internet mediante datos móviles o red Wi-Fi."),
       ("Espacio", "Aproximadamente 100 MB libres de almacenamiento."),
       ("Cuenta", "Correo electrónico y contraseña asignados por el Comisariado (solo socios y contadores)."),
       ("Tienda", "Cuenta de Google activa para descargar la aplicación desde Google Play.")],
      [4, 10.5], "Requisitos mínimos de uso.")

S = "/app/manual_pasos/"


def paso_img(n, texto, arch, titulo, alto=10.0):
    num = fig_n[0] + 1
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_after = Pt(4)
    rr = par.add_run(f"Paso {n}.  "); rr.bold = True; rr.font.color.rgb = RGBColor.from_string(AZUL)
    par.add_run(texto.replace("{f}", f"Figura {num}"))
    figuras([(arch, titulo)], alto)


def pasos_img(lista):
    for i, (t, a, c) in enumerate(lista, 1):
        paso_img(i, t, a, c)


# ======================= 3. INSTALACIÓN =======================
H1("3. Instalación")
p("La aplicación se distribuye a través de la tienda oficial Google Play con el nombre “Gestión de Pozos "
  "Comunitarios”. Para instalarla realice lo siguiente:")
pasos_img([
    ("Abra la aplicación Google Play Store, toque la barra de búsqueda y escriba “Gestión de Pozos Comunitarios”. "
     "La aplicación aparecerá en primer lugar de los resultados, publicada por Gestión Agrícola Escolar, como se "
     "observa en la {f}.", "/app/play_381cc52b1b7f2e86.png", "Búsqueda de la aplicación en Google Play."),
    ("Toque la aplicación para abrir su ficha, donde se muestran el tamaño, la clasificación “Apto para todo "
     "público” y las capturas de pantalla. Presione el botón Instalar ({f}).",
     "/app/play_40182583520d070b.png", "Ficha de la aplicación en Google Play."),
    ("Espere a que termine la descarga; el avance se muestra en porcentaje debajo del nombre de la aplicación "
     "({f}).", "/app/play_a98b372ce1eac741.png", "Descarga e instalación de la aplicación."),
    ("Cuando aparezca la etiqueta “Instalado”, presione Abrir para iniciar la aplicación ({f}). También puede "
     "abrirla desde el ícono que se agrega en la pantalla de su teléfono.",
     "/app/play_85d6d6591112f00c.png", "Aplicación instalada y lista para abrirse."),
])

# ======================= 4. PERFILES =======================
H1("4. Perfiles de usuario")
p("La aplicación cuenta con tres formas de acceso; cada una muestra únicamente las funciones que le "
  "corresponden, como se describe en la Tabla 2:")
tabla(["Perfil", "¿Necesita cuenta?", "Funciones principales"],
      [("Usuario en general", "No", "Ver la lista de socios del pozo, quién está en turno, horas de sobra y llamar a un socio."),
       ("Socio", "Sí", "Consultar su turno y calendario, reportar emergencias, ver notificaciones, multas, estadísticas, historial y registrar agua de sobra."),
       ("Contador", "Sí", "Registrar, editar, cobrar y eliminar multas; registrar recorridos (días sin servicio) del pozo.")],
      [3.5, 3, 8], "Perfiles de usuario de la aplicación.")

# ======================= 5. ACCESO =======================
H1("5. Acceso a la aplicación")
H2("5.1 Pantalla de inicio y selección de pozo")
p("Al abrir la aplicación se muestra la pantalla de bienvenida “Turnos de Pozo”. Cada pozo se identifica con "
  "un color que se mantendrá en todas las pantallas.")
FIG_PUBLICO = fig_n[0] + 3
pasos_img([
    ("Toque el nombre de su pozo (San Isidro, Zapata o Cárdenas); quedará marcado con una palomita, como se "
     "muestra en la {f}.", S + "a1_seleccion.png", "Selección del pozo en la pantalla de inicio."),
    ("Si es socio o contador, presione Iniciar sesión; se abrirá la pantalla de acceso de la {f}.",
     S + "a2_login.png", "Pantalla de inicio de sesión."),
    ("Si solo desea consultar quién tiene el turno, presione Soy usuario (ver turnos); se mostrará la consulta "
     "pública de la {f}.", S + "a3_publico.png", "Consulta pública de socios y turno actual."),
])
H2("5.2 Consulta pública de turnos")
p(f"En la consulta pública (Figura {FIG_PUBLICO}) los socios aparecen en su orden de rotación. El socio que "
  "tiene el turno del día aparece resaltado con la etiqueta “En turno” y los demás como “Inactivo”. Si algún "
  "socio registró horas de agua de sobra, se indica debajo de su nombre. En esta pantalla también puede:")
vineta(" llama directamente al socio.", "Botón de teléfono:")
vineta(" actualiza la información (también puede deslizar la lista hacia abajo).", "Ícono de flechas:")
vineta(" activa o desactiva los avisos del pozo.", "Ícono de campana:")

H2("5.3 Inicio de sesión")
FIG_INICIO = fig_n[0] + 3
pasos_img([
    ("Escriba el correo electrónico que le asignó el Comisariado ({f}).", S + "b1_correo.png",
     "Captura del correo electrónico."),
    ("Escriba su contraseña. Puede tocar el ícono del ojo para comprobar que esté bien escrita ({f}).",
     S + "b2_password.png", "Captura de la contraseña."),
    ("Presione Entrar. La aplicación lo llevará a la pantalla principal de su perfil ({f}).",
     S + "b3_inicio.png", "Pantalla de Inicio del socio."),
])
nota("La sesión permanece abierta en su teléfono; no necesita escribir su contraseña cada vez que abra la aplicación.")

H2("5.4 Recuperar contraseña")
p("Si olvidó su contraseña, presione “Se me olvidó la contraseña” en la pantalla de inicio de sesión. La barra "
  "superior indica la etapa en la que se encuentra:")
pasos_img([
    ("Escriba el correo de su cuenta y presione Siguiente; el código de seguridad se envía al Comisariado ({f}).",
     S + "c1_correo.png", "Solicitud del código de recuperación."),
    ("Solicite el código de 6 dígitos al Comisariado, escríbalo y presione Entrar. El código vence en 15 minutos "
     "({f}).", S + "c2_codigo.png", "Captura del código de seguridad."),
    ("Escriba su nueva contraseña (mínimo 6 caracteres), confírmela y presione Guardar contraseña ({f}).",
     S + "c3_nueva.png", "Captura de la nueva contraseña."),
    ("Se mostrará el mensaje de confirmación; presione Ir a iniciar sesión e ingrese con su nueva contraseña "
     "({f}).", S + "c4_listo.png", "Contraseña actualizada correctamente."),
])

# ======================= 6. SOCIO =======================
salto()
H1("6. Módulo del Socio")
p("Al iniciar sesión como socio se muestra una barra inferior con cuatro secciones: Inicio, Estadísticas, "
  "Historial y Perfil. Desde la pantalla de Inicio también se accede a Emergencia, Notificaciones, Multas y "
  "Agua de sobra.")
H2("6.1 Pantalla de Inicio: turno del día y calendario")
p(f"La pantalla de Inicio (Figura {FIG_INICIO}) muestra la tarjeta “Turno de hoy” con el nombre del socio al que "
  "le corresponde el agua, cuatro botones de acceso rápido y el calendario del mes, en el que se distinguen los "
  "colores descritos en la Tabla 3:")
tabla(["Color en el calendario", "Significado"],
      [("Color del pozo", "Días en los que usted tiene turno."),
       ("Verde", "Turno caído: día sin servicio registrado por el contador."),
       ("Rojo", "Emergencia o día festivo: no hay servicio y la rotación se recorre.")],
      [5, 9.5], "Significado de los colores del calendario.")
H2("6.2 Reportar una emergencia")
pasos_img([
    ("En la pantalla de Inicio presione el botón Emergencia y seleccione el tipo de problema: Fuga en tubería, "
     "Problema en hidrante o Baja presión de agua ({f}).", S + "d2_emergencia.png", "Selección del tipo de emergencia."),
    ("El aviso se envía de inmediato a todos los integrantes del pozo y aparece un mensaje de confirmación en la "
     "parte superior ({f}).", S + "d3_emergencia_ok.png", "Confirmación de emergencia reportada."),
])
H2("6.3 Notificaciones")
pasos_img([
    ("Presione el botón Notificaciones o el ícono de campana para ver los avisos del pozo con su fecha y hora "
     "({f}).", S + "d4_notificaciones.png", "Ventana de notificaciones."),
    ("Presione “Marcar todas como leídas” para limpiar los avisos nuevos ({f}).", S + "d5_leidas.png",
     "Notificaciones marcadas como leídas."),
])
H2("6.4 Consultar multas del pozo")
pasos_img([
    ("En la pantalla de Inicio presione Multas; se mostrará la lista de socios con el número de multas de cada "
     "uno ({f}).", S + "d6_multas.png", "Lista de multas por socio."),
    ("Toque el nombre de un socio para desplegar sus multas con el motivo, el monto y el estado Pagado o No "
     "pagado ({f}).", S + "d7_multas_socio.png", "Multas desplegadas de un socio."),
    ("Toque una multa para abrir su detalle con la descripción, el monto y la fecha ({f}).",
     S + "d8_multa_detalle.png", "Detalle de una multa."),
])
nota("Las multas solo pueden ser registradas y marcadas como pagadas por el contador. El pago se realiza "
     "directamente con el Comisariado.")
H2("6.5 Registrar agua de sobra")
pasos_img([
    ("En la pantalla de Inicio presione Agua de sobra, toque la lista “Selecciona las horas” y elija el número "
     "de horas que no utilizará ({f}).", S + "d9_horas.png", "Selección de horas de sobra."),
    ("Confirme presionando Seguro, o presione Cancelar para regresar ({f}).", S + "d10_seguro.png",
     "Confirmación de horas de sobra."),
    ("Aparecerá el mensaje Listo; presione OK. Las horas se mostrarán en la consulta pública y se notificará al "
     "pozo ({f}).", S + "d11_listo.png", "Horas de sobra registradas."),
    ("Para retirar horas que ya no tiene disponibles, use la lista “Quitar horas de sobra”, elija las horas y "
     "confirme ({f}).", S + "d12_quitar.png", "Retiro de horas de sobra."),
])
H2("6.6 Estadísticas")
pasos_img([
    ("Toque Estadísticas en la barra inferior para consultar los turnos del año, sus multas, el monto pendiente "
     "y pagado, la gráfica de multas pagadas y no pagadas y los socios con más multas ({f}).",
     S + "d13_estadisticas.png", "Pantalla de estadísticas."),
])
H2("6.7 Historial de turnos y exportación a PDF")
pasos_img([
    ("Toque Historial en la barra inferior y presione el selector de mes para elegir el periodo ({f}).",
     S + "d14_meses.png", "Selección del mes del historial."),
    ("Revise la lista de días con su estado (Turno asignado, Festivo o Sin servicio) y presione Exportar a PDF "
     "para guardar, imprimir o compartir sus turnos ({f}).", S + "d15_historial.png", "Historial mensual de turnos."),
])
H2("6.8 Perfil")
pasos_img([
    ("Toque Perfil en la barra inferior para ver su nombre, correo y tipo de usuario ({f}).", S + "d16_perfil.png",
     "Pantalla de perfil."),
    ("Active el interruptor Modo oscuro para cambiar los colores de la aplicación a tonos oscuros ({f}).",
     S + "d17_oscuro.png", "Perfil con modo oscuro activado."),
    ("Para cambiar su contraseña, escriba la actual y la nueva (mínimo 6 caracteres) y presione Actualizar "
     "contraseña ({f}).", S + "d18_cambiar_pass.png", "Cambio de contraseña."),
    ("Presione Cerrar sesión para salir; la aplicación regresará a la pantalla de inicio ({f}).",
     S + "d19_logout.png", "Pantalla de inicio después de cerrar sesión."),
])

# ======================= 7. CONTADOR =======================
salto()
H1("7. Módulo del Contador")
p("El contador de cada pozo dispone de tres secciones en la barra inferior: Multas, Recorrido y Perfil. "
  "La sección Perfil funciona igual que la del socio (apartado 6.8).")
H2("7.1 Registrar una multa")
pasos_img([
    ("Al iniciar sesión se muestra la lista de multas del pozo con el socio, el monto, la fecha y el estado "
     "({f}).", S + "e1_multas.png", "Lista de multas del pozo (contador)."),
    ("Presione el botón Nueva multa, ubicado en la esquina inferior derecha; se abrirá el formulario ({f}).",
     S + "e2_form.png", "Formulario de nueva multa."),
    ("Toque el campo Socio y seleccione al socio sancionado de la lista ({f}).", S + "e3_socios.png",
     "Selección del socio."),
    ("Escriba la Descripción o motivo de la multa ({f}).", S + "e4_descripcion.png", "Captura de la descripción."),
    ("Escriba el Monto en pesos y verifique la Fecha (AAAA-MM-DD), que por defecto es la del día ({f}).",
     S + "e5_monto.png", "Captura del monto y la fecha."),
    ("Presione Crear multa. La multa se agrega a la lista y el socio recibe una notificación ({f}).",
     S + "e6_creada.png", "Multa registrada en la lista."),
])
H2("7.2 Marcar como pagada, editar o eliminar una multa")
pasos_img([
    ("Toque una multa para desplegar su descripción y los botones de acción ({f}).", S + "e7_detalle.png",
     "Multa desplegada con sus acciones."),
    ("Presione Marcar pagada para registrar el pago; el estado cambia a Pagado y se guarda la fecha de pago "
     "({f}).", S + "e8_pagada.png", "Registro del pago de una multa."),
    ("Presione Editar para corregir el socio, la descripción o el monto y después presione Guardar cambios "
     "({f}).", S + "e9_editar.png", "Edición de una multa."),
    ("Presione el ícono de bote de basura para eliminar la multa y confirme con “Sí, eliminar” ({f}).",
     S + "e10_eliminar.png", "Confirmación para eliminar una multa."),
])
nota("Elimine una multa solo cuando haya sido registrada por error. Si la multa ya se cubrió, utilice "
     "“Marcar pagada” para conservar el registro como evidencia.", titulo="Importante:", color="FDE2E1")
H2("7.3 Registrar un recorrido (día sin servicio)")
p("Cuando un día no habrá servicio de agua, el contador lo registra para que la rotación se recorra "
  "automáticamente y ningún socio pierda su turno:")
pasos_img([
    ("Toque Recorrido en la barra inferior para ver el calendario del mes ({f}).", S + "e11_recorrido.png",
     "Calendario de recorridos."),
    ("Toque el día sin servicio y elija el motivo: Turno caído o Emergencia ({f}).", S + "e12_motivo.png",
     "Selección del motivo del recorrido."),
    ("Si eligió Emergencia, escriba la razón en el cuadro de texto ({f}).", S + "e13_razon.png",
     "Captura de la razón de la emergencia."),
    ("Presione Confirmar. El día se pinta en el calendario y se notifica a los socios ({f}).",
     S + "e14_confirmado.png", "Recorrido registrado."),
    ("Para deshacerlo, toque de nuevo el día marcado y presione Quitar recorrido ({f}).", S + "e15_quitar.png",
     "Quitar un recorrido registrado."),
])


# ======================= 8. ROTACIÓN =======================
H1("8. ¿Cómo se asignan los turnos?")
p("La aplicación calcula el turno de forma automática, sin intervención manual, siguiendo las reglas acordadas "
  "por la asamblea:")
vineta(" cada socio tiene un número de orden; los turnos avanzan uno por uno y al llegar al último se reinicia con el primero.", "Rotación en orden:")
vineta(" el turno cambia todos los días a las 18:00 horas.", "Hora de corte:")
vineta(" en días festivos y días sin servicio no se asigna turno; el socio que seguía conserva su lugar para el siguiente día válido.", "Días no válidos:")
vineta(" todos los usuarios ven la misma información al mismo tiempo, lo que evita confusiones y reclamos.", "Transparencia:")
nota("Ejemplo: si al socio número 5 le correspondía el martes y ese día se registra un recorrido, el turno del "
     "socio 5 pasa al miércoles, y así sucesivamente para los demás.", titulo="Ejemplo:", color="E3F2E1")

# ======================= 9. PROBLEMAS =======================
H1("9. Preguntas frecuentes y solución de problemas")
p("En la Tabla 4 se presentan las situaciones más comunes que pueden presentarse al usar la aplicación, su posible causa y la solución recomendada.")
tabla(["Situación", "Posible causa", "Solución"],
      [("No puedo iniciar sesión", "Correo o contraseña incorrectos.", "Verifique los datos con el ícono del ojo o use “Se me olvidó la contraseña”."),
       ("No llega el código de recuperación", "El código se envía al Comisariado, no al socio.", "Solicítelo directamente al Comisariado; recuerde que vence en 15 minutos."),
       ("La información no se actualiza", "Sin conexión a internet.", "Revise sus datos móviles o Wi-Fi y deslice la pantalla hacia abajo para actualizar."),
       ("No aparece mi turno", "El día es festivo o se registró un recorrido.", "Revise los colores del calendario; el turno se recorre al siguiente día válido."),
       ("No recibo avisos", "Los avisos del pozo están desactivados.", "En la consulta pública toque el ícono de campana para activarlos y acepte el permiso de notificaciones."),
       ("Veo una multa que no reconozco", "Registro realizado por el contador.", "Acuda con el contador o el Comisariado para aclararla.")],
      [4, 4.5, 6], "Problemas frecuentes y su solución.")

# ======================= 10. GLOSARIO =======================
H1("10. Glosario")
p("En la Tabla 5 se definen los términos utilizados con mayor frecuencia en este manual.")
tabla(["Término", "Definición"],
      [("Turno", "Día asignado a un socio para utilizar el agua del pozo."),
       ("Rotación", "Orden en que los socios reciben su turno, de forma sucesiva y cíclica."),
       ("Recorrido", "Día sin servicio que hace que la rotación se recorra un día."),
       ("Turno caído", "Recorrido ocasionado por una falla o suspensión del servicio."),
       ("Agua de sobra", "Horas de turno que un socio no utilizará y pone a disposición de los demás."),
       ("Contador", "Encargado del pozo que registra multas y recorridos."),
       ("Comisariado", "Órgano del ejido que administra y representa a los pozos comunitarios.")],
      [4, 10.5], "Glosario de términos.")

# ======================= 11. SOPORTE =======================
H1("11. Soporte y contacto")
p("Para dudas sobre turnos, multas o cuentas de usuario, acuda al Subcomité Técnico de Aguas y Pozos "
  "Hidráulicos del Comisariado Ejidal de San José Carpinteros. Para reportar fallas de la aplicación, "
  "comuníquese con el responsable del sistema, cuyos datos se muestran en la Tabla 6:")
tabla(["Concepto", "Dato"],
      [("Responsable del sistema", "Juan José Huerta Rojas"),
       ("Correo de contacto", "______________________________"),
       ("Teléfono", "______________________________"),
       ("Horario de atención", "______________________________")],
      [5, 9.5], "Datos de contacto para soporte.")


# ======================= ESTILOS DE TÍTULOS E ÍNDICES =======================
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.text.paragraph import Paragraph


def estilo_cap(nombre):
    s = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = doc.styles["Caption"]
    s.font.name = "Arial"; s.font.size = Pt(9); s.font.italic = False; s.font.bold = False
    s.font.color.rgb = RGBColor(0, 0, 0)
    return s


cf, ct = estilo_cap("Cap_Figura"), estilo_cap("Cap_Tabla")
figs, tabs = [], []
for el in doc.element.body.iter(qn("w:p")):
    par = Paragraph(el, doc)
    t = par.text.strip()
    if par.style.name == "Caption" and t.startswith("Figura "):
        par.style = cf; figs.append(t)
    elif par.style.name == "Caption" and t.startswith("Tabla "):
        par.style = ct; tabs.append(t)

ind = doc.styles.add_style("Indice_Figuras_Tablas", WD_STYLE_TYPE.PARAGRAPH)
ind.base_style = doc.styles["Normal"]
ind.paragraph_format.line_spacing = 1.15
ind.paragraph_format.space_after = Pt(3)
ind.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def r_fld(tipo):
    r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tipo); r.append(f); return r


def r_txt(t, tab=False):
    r = OxmlElement("w:r")
    if tab:
        r.append(OxmlElement("w:tab"))
    x = OxmlElement("w:t"); x.set(qn("xml:space"), "preserve"); x.text = t; r.append(x); return r


def bloque_indice(titulo, entradas, estilo):
    br = OxmlElement("w:p"); rr = OxmlElement("w:r"); b = OxmlElement("w:br"); b.set(qn("w:type"), "page")
    rr.append(b); br.append(rr)
    tp = OxmlElement("w:p"); pPr = OxmlElement("w:pPr"); jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center")
    pPr.append(jc); tp.append(pPr)
    rt = r_txt(titulo); rPr = OxmlElement("w:rPr"); rPr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "28"); rPr.append(sz); rt.insert(0, rPr); tp.append(rt)
    els = [br, tp]
    for k, e in enumerate(entradas):
        p_ = OxmlElement("w:p"); pPr = OxmlElement("w:pPr"); ps_ = OxmlElement("w:pStyle")
        ps_.set(qn("w:val"), ind.style_id); pPr.append(ps_); p_.append(pPr)
        if k == 0:
            p_.append(r_fld("begin"))
            r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
            it.text = f' TOC \\h \\z \\t "{estilo},1" '; r.append(it); p_.append(r)
            p_.append(r_fld("separate"))
        p_.append(r_txt(e)); p_.append(r_txt(str(PAGINAS.get(e, "")), tab=True))
        if k == len(entradas) - 1:
            p_.append(r_fld("end"))
        els.append(p_)
    return els


import json, sys
PAGINAS = json.load(open(sys.argv[1])) if len(sys.argv) > 1 else {}
pos = pt._p
for el in bloque_indice("ÍNDICE DE FIGURAS", figs, "Cap_Figura") + bloque_indice("ÍNDICE DE TABLAS", tabs, "Cap_Tabla"):
    pos.addnext(el); pos = el
json.dump({"figs": figs, "tabs": tabs}, open("/tmp/caps2.json", "w"), ensure_ascii=False)

# Actualizar índice al abrir
st_el = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
sig = next((st_el.find(qn(t)) for t in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat",
                                          "w:docVars", "w:rsids") if st_el.find(qn(t)) is not None), None)
sig.addprevious(uf) if sig is not None else st_el.append(uf)

doc.save("/app/Manual_de_Usuario_Pasos.docx")
print("ok", fig_n[0], "figuras", tab_n[0], "tablas")
