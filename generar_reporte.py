# -*- coding: utf-8 -*-
"""Genera el Reporte Final de Residencia Profesional ya formateado (ITS Tepeaca)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

doc = Document()

# ---- Estilo base: Arial 11 ----
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
pf.space_after = Pt(0)
pf.space_before = Pt(0)

# ---- Margenes: sup 2.5, inf 2.5, izq 3, der 2.5 ----
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)


def _set_font(run, name="Arial", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def cap_titulo(texto):
    """Titulo de capitulo: Arial 14, mayusculas, centrado, negrita."""
    doc.add_page_break()
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto.upper())
    _set_font(r, size=14, bold=True)
    doc.add_page_break()


def h1(texto):
    """Titulo tema 1: Arial 14, negritas, mayusculas, izquierda."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(texto.upper())
    _set_font(r, size=14, bold=True)
    return p


def h2(texto):
    """Subtitulo (Titulo 2): Arial 12, negritas, mayusculas, izquierda."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(texto.upper())
    _set_font(r, size=12, bold=True)
    return p


def h3(texto):
    """Subtema (Titulo 3): Arial 12, negritas, mayus/minus, izquierda."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(texto)
    _set_font(r, size=12, bold=True)
    return p


def parr(texto, justificado=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justificado else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(texto)
    _set_font(r)
    return p


def vinieta(texto):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(texto)
    _set_font(r)
    return p


def centrado(texto, size=11, bold=False, upper=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto.upper() if upper else texto)
    _set_font(r, size=size, bold=bold)
    return p


# =========================================================
# PORTADA
# =========================================================
centrado("INSTITUTO TECNOLÓGICO SUPERIOR DE TEPEACA", size=14, bold=True, upper=True)
doc.add_paragraph()
centrado("[ Insertar logo SEP y logo TecNM ]", size=10)
for _ in range(3):
    doc.add_paragraph()
centrado("TESIS", size=16, bold=True)
doc.add_paragraph()
centrado("“GESTIÓN DE POZOS COMUNITARIOS: APLICACIÓN MÓVIL Y WEB PARA LA "
         "ADMINISTRACIÓN DE TURNOS DE AGUA”", size=13, bold=True)
for _ in range(2):
    doc.add_paragraph()
centrado("QUE PARA OBTENER EL TÍTULO DE:", size=11)
centrado("[ NOMBRE DE LA CARRERA ]", size=12, bold=True, upper=True)
for _ in range(2):
    doc.add_paragraph()
centrado("PRESENTA:", size=11)
centrado("[ NOMBRE(S) Y APELLIDOS DEL ALUMNO, SIN ABREVIATURAS ]", size=12, bold=True)
for _ in range(2):
    doc.add_paragraph()
centrado("ASESOR INTERNO:", size=11)
centrado("[ GRADO ACADÉMICO, NOMBRE Y APELLIDOS ]", size=11, bold=True)
doc.add_paragraph()
centrado("ASESOR EXTERNO:", size=11)
centrado("[ GRADO ACADÉMICO, NOMBRE Y APELLIDOS ]", size=11, bold=True)
for _ in range(3):
    doc.add_paragraph()
centrado("TEPEACA DE NEGRETE, PUE. [ MES ] 2026", size=11, bold=True)

# =========================================================
# AGRADECIMIENTOS
# =========================================================
doc.add_page_break()
h1("Agradecimientos")
parr("Se expresa un profundo agradecimiento a la institución y a los docentes del "
     "Instituto Tecnológico Superior de Tepeaca, cuya guía y conocimientos hicieron "
     "posible la realización del presente trabajo. De igual manera, se reconoce el "
     "apoyo de las comunidades de los pozos San Isidro, Zapata y Cárdenas, quienes "
     "aportaron la información y la confianza necesarias para el desarrollo del "
     "proyecto.")
parr("[ Espacio reservado para que el alumno agregue su dedicatoria personal a "
     "familiares, asesores y compañeros. ]")

# =========================================================
# RESUMEN
# =========================================================
doc.add_page_break()
h1("Resumen")
parr("El presente trabajo describe el diseño y desarrollo de una aplicación híbrida, "
     "móvil y web, para la administración de los turnos de agua de tres pozos "
     "comunitarios: San Isidro, Zapata y Cárdenas. El sistema resuelve la confusión y "
     "los conflictos que ocasionaba el control manual de los turnos, las multas y los "
     "avisos dentro de cada comunidad. La solución fue construida con React Native y "
     "Expo para el cliente, FastAPI para el servidor y MongoDB como base de datos, "
     "empleando autenticación por tokens JWT y una arquitectura de servicios REST. La "
     "aplicación calcula de forma automática la rotación diaria de turnos con corte a "
     "las 18:00 horas, contempla días festivos y días sin servicio, distingue los "
     "perfiles de socio y contador, y permite la gestión de multas y el registro de "
     "agua de sobra. Como resultado se obtuvo un sistema funcional, publicado y "
     "operando de manera independiente, que aporta transparencia y orden a la "
     "distribución del agua en las comunidades participantes.")

# =========================================================
# INDICE GENERAL (marcador)
# =========================================================
doc.add_page_break()
h1("Índice General")
parr("[ Generar el índice automático en Word: menú Referencias → Tabla de contenido. "
     "Aplicar los estilos de título para que se actualice solo. A continuación deben "
     "colocarse también el Índice de Figuras y el Índice de Tablas. ]")

# =========================================================
# CAPITULO I
# =========================================================
cap_titulo("Capítulo I. Generalidades del Proyecto")

h1("1.1 Introducción")
parr("El agua es un recurso indispensable para las comunidades rurales, y su "
     "distribución equitativa representa uno de los mayores retos de organización "
     "social. En muchas localidades el acceso al agua se obtiene mediante pozos "
     "administrados por los propios habitantes, quienes se organizan en turnos para "
     "aprovechar el recurso de manera ordenada. Sin embargo, cuando este control se "
     "realiza de forma manual, mediante libretas o de palabra, surgen confusiones, "
     "olvidos y conflictos entre los socios.")
parr("El presente documento describe el desarrollo de una aplicación híbrida, "
     "disponible para dispositivos móviles Android y para navegadores web, orientada a "
     "la administración de los turnos de agua de tres pozos comunitarios. El trabajo "
     "expone la problemática detectada, los objetivos planteados, el marco teórico que "
     "sustenta las tecnologías utilizadas, la metodología de desarrollo seguida y los "
     "resultados obtenidos. La finalidad es mostrar cómo una herramienta tecnológica "
     "sencilla y accesible puede aportar orden, transparencia y confianza a la "
     "distribución de un recurso tan valioso como el agua.")

h1("1.2 Descripción de la organización y del área de trabajo")
h3("1.2.1 Antecedentes de la organización")
parr("La organización está conformada por tres comités comunitarios encargados de la "
     "administración de los pozos de agua conocidos como San Isidro, Zapata y "
     "Cárdenas. Cada comité agrupa a un conjunto de socios que tienen derecho al uso "
     "del agua y que participan en la rotación de turnos. La administración de cada "
     "pozo recae en la figura del contador, quien lleva el control de la rotación, las "
     "multas y los días sin servicio.")
parr("Estas organizaciones son de tipo comunitario y sin fines de lucro. Su objetivo "
     "principal es garantizar el acceso justo y ordenado al agua para todos sus "
     "integrantes. Aunque no cuentan con una estructura empresarial formal, sí "
     "mantienen una organización interna basada en la confianza y en acuerdos "
     "colectivos, en la que el contador funge como responsable operativo y los socios "
     "como beneficiarios del servicio.")
parr("Misión: administrar de manera justa, ordenada y transparente el uso del agua de "
     "los pozos comunitarios, garantizando que cada socio reciba su turno "
     "correspondiente.")
parr("Visión: consolidar un sistema de organización confiable y moderno que reduzca "
     "los conflictos y facilite la gestión del agua en beneficio de toda la "
     "comunidad.")

h3("1.2.2 Caracterización del área donde se generó el proyecto")
parr("El proyecto se generó en el área de administración de los pozos, específicamente "
     "en las tareas que realiza el contador para llevar el control de los turnos. "
     "Dicha área tiene como objetivo principal registrar y comunicar a cada socio el "
     "día que le corresponde el uso del agua, así como administrar las multas por "
     "incumplimiento y los periodos en los que no hay servicio.")
parr("Esta área se relaciona directamente con todos los socios del pozo, pues es la "
     "encargada de mantener actualizada la información de los turnos y de resolver las "
     "dudas o inconformidades. En cada pozo colabora principalmente el contador, "
     "apoyado por los socios que reportan incidencias o emergencias.")

h1("1.3 Problemas a resolver")
parr("El control de los turnos de agua se realizaba de forma manual, lo que generaba "
     "diversos problemas que afectaban la convivencia y la organización de las "
     "comunidades. A continuación se describen los principales problemas detectados, "
     "ordenados según su prioridad:")
vinieta("Confusión en los turnos: al llevar el control en libretas o de memoria, con "
        "frecuencia dos socios creían tener el mismo turno, o alguno perdía el suyo "
        "por olvido, lo que ocasionaba discusiones.")
vinieta("Falta de transparencia: los socios no tenían una forma clara de consultar "
        "cuándo les tocaba su turno, ni de verificar el orden de la rotación, lo que "
        "generaba desconfianza hacia la administración.")
vinieta("Dificultad para registrar multas y días sin servicio: no existía un registro "
        "ordenado de las multas por incumplimiento ni de los días festivos o de "
        "emergencia en que el pozo no daba servicio.")
vinieta("Comunicación deficiente: los avisos importantes se transmitían de boca en "
        "boca, por lo que muchos socios no se enteraban a tiempo de los cambios en la "
        "rotación.")
parr("En síntesis, la ausencia de una herramienta que centralizara y automatizara la "
     "información provocaba desorden, pérdida de tiempo y conflictos constantes entre "
     "los integrantes de las comunidades.")

h1("1.4 Objetivos")
h3("1.4.1 Objetivo General")
parr("Desarrollar una aplicación híbrida, para dispositivos móviles y web, que "
     "administre de forma automática los turnos de agua de tres pozos comunitarios, "
     "empleando las tecnologías React Native, Expo, FastAPI y MongoDB, con el fin de "
     "aportar orden y transparencia a la distribución del recurso.")
h3("1.4.2 Objetivos Específicos")
vinieta("Analizar la forma en que se administran actualmente los turnos de agua en "
        "los pozos comunitarios para determinar los requerimientos del sistema.")
vinieta("Diseñar la base de datos y las interfaces de usuario diferenciadas para los "
        "perfiles de socio y contador.")
vinieta("Implementar un mecanismo automático de rotación diaria de turnos con corte a "
        "las 18:00 horas, que contemple días festivos y días sin servicio.")
vinieta("Desarrollar los módulos de gestión de multas, registro de agua de sobra y "
        "avisos de la comunidad.")
vinieta("Incorporar un sistema de autenticación seguro y de recuperación de "
        "contraseña por correo electrónico.")
vinieta("Publicar la aplicación de manera que pueda ser instalada por los socios en "
        "sus dispositivos móviles.")

h1("1.5 Justificación de la propuesta")
parr("La propuesta se justifica por la necesidad de contar con una herramienta que "
     "elimine la confusión y los conflictos derivados del control manual de los turnos "
     "de agua. Al automatizar la rotación y centralizar la información, se garantiza "
     "que cada socio conozca con certeza el día que le corresponde el uso del agua, lo "
     "que fortalece la confianza dentro de la comunidad.")
parr("Desde el punto de vista social, el proyecto beneficia directamente a los socios "
     "de los tres pozos, quienes obtienen un acceso más justo y ordenado al recurso. "
     "Desde el punto de vista práctico, el contador reduce su carga de trabajo al "
     "contar con un sistema que realiza automáticamente los cálculos de la rotación y "
     "el registro de multas.")
parr("Además, el uso de tecnologías modernas y de código mantenible permite que la "
     "solución sea escalable y pueda replicarse en otras comunidades con "
     "necesidades similares. La factibilidad del proyecto se sustenta en el empleo de "
     "herramientas de amplia difusión y bajo costo, así como en el aprovechamiento de "
     "servicios en la nube para su operación continua.")

# =========================================================
# CAPITULO II
# =========================================================
cap_titulo("Capítulo II. Marco Teórico")

h1("2.1 Fundamentos teóricos")
parr("En este apartado se presentan los conceptos y tecnologías que sustentan el "
     "desarrollo de la aplicación. El conjunto de fundamentos teóricos permite "
     "comprender las decisiones técnicas adoptadas durante el proyecto.")

h3("Aplicaciones híbridas")
parr("Una aplicación híbrida es aquella que, a partir de una sola base de código, "
     "puede ejecutarse en distintas plataformas, como teléfonos móviles y navegadores "
     "web. Este enfoque reduce el tiempo y el costo de desarrollo, ya que evita "
     "programar por separado cada versión de la aplicación.")

h3("React Native y Expo")
parr("React Native es un marco de trabajo de código abierto que permite construir "
     "aplicaciones móviles utilizando el lenguaje JavaScript y la biblioteca React. "
     "Su principal ventaja es que genera componentes nativos, lo que ofrece un "
     "rendimiento cercano al de las aplicaciones desarrolladas de forma nativa. Expo "
     "es un conjunto de herramientas y servicios construido sobre React Native que "
     "simplifica el desarrollo, las pruebas y la generación de las versiones "
     "instalables de la aplicación.")

h3("FastAPI")
parr("FastAPI es un marco de trabajo moderno del lenguaje Python, orientado a la "
     "creación de interfaces de programación de aplicaciones (API). Se caracteriza por "
     "su alto rendimiento, su facilidad de uso y su validación automática de datos. "
     "En el presente proyecto se utilizó para construir el servidor que atiende las "
     "peticiones de la aplicación cliente.")

h3("MongoDB")
parr("MongoDB es un sistema de gestión de bases de datos orientado a documentos, "
     "clasificado dentro de las bases de datos NoSQL. En lugar de almacenar la "
     "información en tablas, la organiza en documentos con formato flexible similar al "
     "de la notación JSON. Esta característica facilita el manejo de estructuras de "
     "datos variables, como las utilizadas para representar los pozos, los usuarios y "
     "los turnos.")

h3("Arquitectura REST y autenticación JWT")
parr("La transferencia de estado representacional (REST) es un estilo de arquitectura "
     "para el diseño de servicios web, en el que la comunicación entre el cliente y el "
     "servidor se realiza mediante peticiones sobre el protocolo HTTP. Para proteger "
     "el acceso a la información se empleó el estándar JSON Web Token (JWT), que "
     "consiste en un token firmado digitalmente que valida la identidad del usuario en "
     "cada petición sin necesidad de mantener una sesión en el servidor.")

h3("Servicios en la nube")
parr("Los servicios en la nube permiten alojar aplicaciones y bases de datos en "
     "servidores accesibles a través de Internet. En este proyecto se utilizaron "
     "servicios de despliegue para el servidor y de almacenamiento gestionado para la "
     "base de datos, lo que garantiza la disponibilidad continua del sistema sin "
     "requerir infraestructura física propia.")

# =========================================================
# CAPITULO III
# =========================================================
cap_titulo("Capítulo III. Metodología")

h1("3.1 Análisis y Requerimientos")
parr("La primera etapa del proyecto consistió en analizar la forma en que se "
     "administraban los turnos de agua en los pozos comunitarios. Para ello se "
     "recabó información directamente con los contadores y socios, se clasificó de "
     "acuerdo con las tareas que realizaban y se determinaron los requerimientos "
     "funcionales del sistema.")
parr("Como resultado del análisis se identificaron los siguientes requerimientos "
     "funcionales:")
vinieta("El sistema debe permitir el inicio de sesión diferenciado para los perfiles "
        "de socio y contador.")
vinieta("El sistema debe calcular automáticamente la rotación diaria de turnos, con "
        "un corte a las 18:00 horas.")
vinieta("El sistema debe contemplar los días festivos y los días sin servicio, "
        "recorriendo la rotación sin saltarse a ningún socio.")
vinieta("El contador debe poder registrar multas, días sin servicio y recorridos.")
vinieta("Los socios deben poder consultar su turno, reportar emergencias y registrar "
        "el agua de sobra.")
vinieta("El sistema debe permitir la recuperación de contraseña mediante el correo "
        "electrónico.")
parr("Entre los requerimientos no funcionales se establecieron la disponibilidad "
     "continua del servicio, la seguridad de la información mediante autenticación por "
     "tokens y una interfaz sencilla y clara, adecuada para usuarios sin experiencia "
     "técnica.")

h1("3.2 Metodología")
parr("Para el desarrollo del sistema se adoptó una metodología de tipo incremental e "
     "iterativa, propia del desarrollo ágil de software. Bajo este enfoque, el "
     "producto se construyó por etapas, entregando funcionalidades operativas de "
     "manera progresiva y sometiéndolas a revisión antes de continuar con la "
     "siguiente. Este método resultó adecuado por permitir la validación constante de "
     "los avances con los usuarios y la corrección oportuna de errores.")
parr("El proceso se organizó en las siguientes fases: análisis de requerimientos, "
     "diseño de la base de datos y de las interfaces, implementación de los módulos, "
     "pruebas de funcionamiento y despliegue de la aplicación.")

h1("3.3 Diseño e Implementación")
h3("Diseño de la base de datos")
parr("La base de datos se modeló empleando colecciones de documentos, de acuerdo con "
     "la naturaleza no relacional de MongoDB. Las colecciones principales son: "
     "«pozos», que almacena los datos de cada pozo y su fecha de inicio de rotación; "
     "«usuarios», que guarda la información de socios y contadores junto con su rol y "
     "orden en la rotación; «multas», que registra las sanciones aplicadas; "
     "«dias_sin_servicio», que almacena los recorridos y días sin servicio; y "
     "«emergencias», que guarda los reportes realizados por los socios.")
h3("Diseño de las interfaces")
parr("Se diseñaron interfaces diferenciadas para cada perfil. La interfaz del socio "
     "muestra un calendario que resalta con colores su turno, los días festivos y los "
     "días sin servicio, además de las secciones de multas, agua de sobra y avisos. "
     "La interfaz del contador incorpora las herramientas de administración, entre "
     "ellas el registro de multas y la gestión de recorridos, que permiten marcar un "
     "día como turno caído o como emergencia.")
h3("Implementación")
parr("La implementación se realizó separando el sistema en dos componentes: el "
     "cliente, desarrollado con React Native y Expo, y el servidor, desarrollado con "
     "FastAPI. La comunicación entre ambos se efectúa mediante servicios REST "
     "protegidos con tokens JWT. La lógica de rotación de turnos se programó en el "
     "servidor, evaluando la fecha y la hora locales (zona horaria del centro de "
     "México, UTC-6) para determinar el socio en turno, recorriendo la rotación cuando "
     "se detecta un día festivo o un día sin servicio.")
parr("El servidor se desplegó en un servicio en la nube y la base de datos se alojó en "
     "un servicio de almacenamiento gestionado. La aplicación cliente se compiló en "
     "formato instalable para dispositivos Android y se distribuyó a través de la "
     "plataforma de aplicaciones correspondiente, quedando disponible para su "
     "instalación por parte de los socios.")

# =========================================================
# CAPITULO IV
# =========================================================
cap_titulo("Capítulo IV. Resultados")

h1("4.1 Resultados")
parr("Como resultado del proyecto se obtuvo una aplicación híbrida completamente "
     "funcional, capaz de administrar los turnos de agua de los tres pozos "
     "comunitarios de manera automática y transparente. A continuación se describen "
     "los principales resultados obtenidos, tanto de forma cualitativa como "
     "cuantitativa.")

h3("Funcionamiento del sistema")
parr("El sistema calcula correctamente la rotación diaria de turnos, realizando el "
     "cambio de turno de manera puntual a las 18:00 horas y respetando los días "
     "festivos y los días sin servicio. Las pruebas realizadas confirmaron que la "
     "rotación avanza sin saltarse a ningún socio cuando se presenta un día no "
     "laborable, lo que resuelve uno de los principales problemas identificados.")

h3("Módulos desarrollados")
vinieta("Autenticación y recuperación de contraseña por correo electrónico, con "
        "acceso diferenciado para socios y contadores.")
vinieta("Calendario de turnos con codificación por colores para turnos, días "
        "festivos, turnos caídos y emergencias.")
vinieta("Gestión de multas por parte del contador.")
vinieta("Registro de recorridos y días sin servicio, exclusivo del contador.")
vinieta("Reporte de emergencias y registro de agua de sobra por parte de los socios.")

h3("Pruebas de funcionamiento")
parr("Se realizaron pruebas de funcionamiento sobre los cálculos de la rotación de "
     "turnos, verificando distintos escenarios: días normales, días festivos, días "
     "sin servicio y cambios de turno alrededor de las 18:00 horas. En todos los "
     "casos el sistema respondió de acuerdo con lo esperado. Asimismo, se probó la "
     "correcta separación de la información entre los distintos pozos, garantizando "
     "que cada comunidad visualice únicamente sus propios datos.")

h3("Pruebas de usabilidad")
parr("Las interfaces resultaron claras y sencillas de utilizar para los socios, "
     "incluso para aquellos sin experiencia previa con aplicaciones móviles. El uso de "
     "colores en el calendario facilitó la identificación inmediata del turno "
     "correspondiente y de los días especiales.")

h3("Despliegue y disponibilidad")
parr("El servidor y la base de datos operan de manera continua en servicios en la "
     "nube, lo que asegura la disponibilidad del sistema durante las veinticuatro "
     "horas del día. La aplicación se compiló en formato instalable y se publicó, "
     "quedando disponible para su descarga e instalación en los dispositivos móviles "
     "de los socios.")

h3("Seguridad")
parr("La información de los usuarios se protege mediante contraseñas cifradas y "
     "autenticación por tokens JWT. La comunicación entre la aplicación y el servidor "
     "se realiza sobre canales cifrados, garantizando la confidencialidad de los "
     "datos.")

parr("[ Nota: en este apartado se recomienda insertar las capturas de pantalla de la "
     "aplicación —pantalla de inicio de sesión, calendario de turnos, gestión de "
     "multas, etc.— con su respectivo título de figura en Arial 9, centrado. ]")

# =========================================================
# CONCLUSIONES
# =========================================================
cap_titulo("Conclusiones")
h1("Conclusiones del Proyecto")
parr("El desarrollo de la aplicación para la gestión de los pozos comunitarios "
     "permitió cumplir con el objetivo general planteado, al obtener un sistema "
     "híbrido, funcional y publicado, que administra de forma automática los turnos de "
     "agua de tres pozos. La solución resolvió la problemática de confusión y "
     "desorganización que existía con el control manual, aportando transparencia y "
     "confianza a las comunidades.")
parr("Se comprobó que el uso de tecnologías modernas, como React Native, Expo, "
     "FastAPI y MongoDB, permite construir soluciones robustas y escalables a partir "
     "de una sola base de código, aprovechable tanto en dispositivos móviles como en "
     "navegadores web. La automatización de la rotación de turnos, considerando el "
     "corte a las 18:00 horas y los días sin servicio, resultó ser el componente de "
     "mayor valor para los usuarios.")
parr("El proyecto demostró ser pertinente y viable, y sienta las bases para su "
     "posible replicación en otras comunidades que enfrenten problemáticas similares "
     "en la distribución de recursos.")

# =========================================================
# RECOMENDACIONES
# =========================================================
cap_titulo("Recomendaciones")
h1("Recomendaciones")
parr("Se recomienda capacitar a los contadores de cada pozo en el uso de las "
     "herramientas de administración, con el fin de aprovechar al máximo las "
     "funcionalidades del sistema. Asimismo, conviene mantener actualizada la "
     "información de los socios y de los días festivos para garantizar la exactitud de "
     "la rotación.")
parr("Como trabajo futuro, se sugiere incorporar reportes estadísticos sobre el uso "
     "del agua, así como evaluar la incorporación de notificaciones automáticas que "
     "alerten a cada socio cuando se acerque su turno. También se recomienda realizar "
     "respaldos periódicos de la base de datos para resguardar la información ante "
     "cualquier eventualidad.")

# =========================================================
# EXPERIENCIA PERSONAL
# =========================================================
cap_titulo("Experiencia Personal Profesional Adquirida")
h1("Experiencia Personal Profesional Adquirida")
parr("Durante el desarrollo del presente proyecto se adquirieron conocimientos y "
     "habilidades de gran valor para el ejercicio profesional. Se profundizó en el "
     "manejo de tecnologías actuales para el desarrollo de aplicaciones móviles y web, "
     "así como en el diseño de bases de datos no relacionales y en la construcción de "
     "servicios web seguros.")
parr("Igualmente, se fortalecieron competencias relacionadas con el análisis de "
     "requerimientos, la comunicación con los usuarios finales y la resolución de "
     "problemas reales de una comunidad. La experiencia de llevar un proyecto desde su "
     "concepción hasta su publicación permitió comprender la importancia de cada etapa "
     "del ciclo de vida del software, así como la satisfacción de aportar una solución "
     "tecnológica útil para las personas.")
parr("[ Espacio reservado para que el alumno complemente con su experiencia personal, "
     "los retos enfrentados y los aprendizajes obtenidos. ]")

# =========================================================
# COMPETENCIAS
# =========================================================
cap_titulo("Competencias")
h1("Competencias Desarrolladas y/o Aplicadas")
vinieta("Análisis y determinación de requerimientos de software a partir de "
        "necesidades reales.")
vinieta("Diseño y modelado de bases de datos orientadas a documentos.")
vinieta("Desarrollo de aplicaciones móviles y web con React Native y Expo.")
vinieta("Construcción de servicios web (API REST) con FastAPI y Python.")
vinieta("Implementación de mecanismos de seguridad mediante autenticación por tokens "
        "JWT.")
vinieta("Despliegue de aplicaciones y bases de datos en servicios en la nube.")
vinieta("Documentación técnica y elaboración de reportes profesionales.")
parr("Estas competencias contribuyen de manera directa al perfil de egreso, al "
     "integrar los conocimientos teóricos con la práctica profesional en el desarrollo "
     "de soluciones informáticas.")

# =========================================================
# FUENTES DE INFORMACION
# =========================================================
cap_titulo("Fuentes de Información")
h1("Referencias Bibliográficas y Virtuales")
parr("[1] Meta Platforms, Inc. (2024). React Native – Documentación oficial. "
     "https://reactnative.dev")
parr("[2] Expo. (2024). Expo Documentation. https://docs.expo.dev")
parr("[3] Ramírez, S. (2024). FastAPI – The modern, fast web framework for building "
     "APIs with Python. https://fastapi.tiangolo.com")
parr("[4] MongoDB, Inc. (2024). MongoDB Manual. https://www.mongodb.com/docs")
parr("[5] Internet Engineering Task Force. (2015). RFC 7519: JSON Web Token (JWT). "
     "https://datatracker.ietf.org/doc/html/rfc7519")
parr("[6] Pressman, R. S. (2010). Ingeniería del software: un enfoque práctico "
     "(7a ed.). México: McGraw-Hill.")
parr("[ Nota: se recomienda administrar las referencias con la herramienta de Word "
     "(menú Referencias → Insertar cita, formato APA Sixth Edition) y citarlas "
     "numéricamente dentro del texto. ]")

# =========================================================
# ANEXOS
# =========================================================
cap_titulo("Anexos")
h1("Anexos")
parr("Anexo A. Carta de autorización de las comunidades para la realización y "
     "titulación del proyecto. [ Adjuntar documento. ]")
parr("Anexo B. Capturas de pantalla de la aplicación en funcionamiento.")
parr("Anexo C. Enlace de descarga de la aplicación y evidencias de publicación.")

# =========================================================
# INDICES
# =========================================================
doc.add_page_break()
h1("Índice de Figuras")
parr("[ Colocar aquí la lista de figuras. Formato: «Figura 1. Descripción de la "
     "imagen.» en Arial 9, según la secuencia. ]")
doc.add_paragraph()
h1("Índice de Tablas")
parr("[ Colocar aquí la lista de tablas. Formato: «Tabla 1. Nombre de la tabla.» "
     "según la secuencia. ]")

doc.save("/app/Reporte_Residencia_Pozos.docx")
print("Documento generado: /app/Reporte_Residencia_Pozos.docx")
