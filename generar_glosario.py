from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

T = {
    "Agua de sobra": "Horas de un turno de riego que el socio no utilizará y que pone a disposición de los demás integrantes del pozo mediante la aplicación.",
    "Android": "Sistema operativo móvil desarrollado por Google, utilizado en la mayoría de los teléfonos inteligentes y plataforma en la que se publicó la aplicación.",
    "API (Interfaz de Programación de Aplicaciones)": "Conjunto de reglas y servicios que permite la comunicación entre la aplicación móvil y el servidor para consultar o registrar información.",
    "API REST": "Estilo de arquitectura para servicios web que utiliza los métodos del protocolo HTTP (GET, POST, PUT, DELETE) para manipular recursos identificados por una dirección URL.",
    "Aritmética modular": "Rama de las matemáticas que trabaja con los residuos de una división; en el proyecto se utiliza para calcular de forma cíclica qué socio tiene el turno de cada día.",
    "Asamblea General de Ejidatarios": "Máxima autoridad del ejido, integrada por todos los ejidatarios, en la que se toman los acuerdos sobre el uso de las tierras y los recursos comunes.",
    "Autenticación": "Proceso mediante el cual el sistema verifica la identidad de un usuario a través de su correo electrónico y contraseña.",
    "Backend": "Parte del sistema que se ejecuta en el servidor y se encarga de la lógica de negocio, la seguridad y el acceso a la base de datos.",
    "Base de datos": "Conjunto organizado de información almacenada de forma digital que puede consultarse, actualizarse y administrarse por medio de un sistema.",
    "Base de datos NoSQL": "Tipo de base de datos que no utiliza tablas relacionales, sino estructuras flexibles como documentos; MongoDB es un ejemplo.",
    "Calendario de rotación": "Vista mensual de la aplicación que muestra los días de turno de cada socio, así como los festivos y los días sin servicio.",
    "Cifrado de contraseñas (hash)": "Técnica de seguridad que transforma la contraseña en un valor irreversible antes de guardarla, de modo que no pueda leerse aunque se acceda a la base de datos.",
    "Cliente-servidor": "Modelo de comunicación en el que un dispositivo (cliente) solicita información o servicios a otro equipo central (servidor) a través de una red.",
    "Colección": "Agrupación de documentos dentro de una base de datos MongoDB, equivalente a una tabla en una base de datos relacional.",
    "Comisariado Ejidal": "Órgano ejecutivo y de representación del ejido, integrado por presidente, secretario y tesorero, responsable de administrar los recursos comunes como los pozos.",
    "Contador": "Perfil de usuario encargado de registrar multas, marcar pagos y registrar los días sin servicio de su pozo.",
    "Corte de las 18:00 horas": "Regla del sistema que establece la hora en que cambia el turno de un socio al siguiente.",
    "Despliegue": "Proceso de instalar y poner en funcionamiento una aplicación o servidor en un entorno accesible para los usuarios finales.",
    "Día sin servicio (recorrido)": "Día en el que no se distribuye agua por falla, mantenimiento o emergencia; la rotación se recorre para que ningún socio pierda su turno.",
    "EAS (Expo Application Services)": "Servicio de Expo que permite compilar la aplicación y generar los archivos de instalación para su publicación en tiendas.",
    "Ejido": "Forma de propiedad social de la tierra en México, reconocida por la Ley Agraria, cuyos integrantes se denominan ejidatarios.",
    "Emergencia": "Reporte de una falla en la red de agua, como fuga en tubería, problema en hidrante o baja presión, que se notifica a todos los integrantes del pozo.",
    "Expo": "Plataforma y conjunto de herramientas que facilita el desarrollo, la prueba y la compilación de aplicaciones móviles creadas con React Native.",
    "FastAPI": "Framework del lenguaje Python utilizado para construir el servidor y los servicios web de la aplicación.",
    "Festivo": "Día no laborable definido en el sistema en el que no se asigna turno y la rotación continúa en el siguiente día válido.",
    "Frontend": "Parte del sistema con la que interactúa directamente el usuario; en este proyecto corresponde a la aplicación móvil.",
    "Google Play Console": "Plataforma de Google para administrar, probar y publicar aplicaciones Android en la tienda Google Play.",
    "HTTPS": "Versión segura del protocolo HTTP que cifra la información transmitida entre la aplicación y el servidor.",
    "Interfaz de usuario": "Conjunto de pantallas, botones y elementos visuales a través de los cuales el usuario interactúa con la aplicación.",
    "JSON": "Formato de texto ligero utilizado para intercambiar datos entre la aplicación móvil y el servidor.",
    "JWT (JSON Web Token)": "Estándar abierto que genera un token firmado digitalmente para identificar al usuario en cada solicitud después de iniciar sesión.",
    "Modelo matemático": "Representación mediante expresiones matemáticas de un proceso real; en el proyecto describe el cálculo del turno diario de cada socio.",
    "MongoDB": "Sistema gestor de bases de datos NoSQL orientado a documentos, utilizado para almacenar la información de pozos, usuarios, multas y emergencias.",
    "MongoDB Atlas": "Servicio en la nube que permite alojar y administrar bases de datos MongoDB sin necesidad de un servidor propio.",
    "Multa": "Sanción económica registrada por el contador a un socio por incumplir los acuerdos del pozo; puede tener estado pagado o no pagado.",
    "Notificación": "Aviso que muestra la aplicación para informar sobre emergencias, multas, recorridos, horas de sobra o cambios de turno.",
    "Pozo comunitario": "Infraestructura hidráulica de uso común que extrae agua del subsuelo para el riego agrícola de los socios.",
    "Prueba cerrada": "Etapa de Google Play en la que solo un grupo de verificadores invitados puede instalar y probar la aplicación antes de su publicación.",
    "Python": "Lenguaje de programación de alto nivel utilizado para desarrollar el servidor de la aplicación.",
    "React Native": "Framework de desarrollo creado por Meta que permite construir aplicaciones móviles nativas utilizando JavaScript.",
    "Render": "Plataforma de alojamiento en la nube en la que se desplegó el servidor de la aplicación.",
    "Requerimiento funcional": "Descripción de una función o servicio que el sistema debe realizar, como calcular el turno o registrar una multa.",
    "Requerimiento no funcional": "Característica de calidad que debe cumplir el sistema, como seguridad, rendimiento, disponibilidad o facilidad de uso.",
    "Rol de usuario": "Tipo de acceso asignado a cada cuenta (socio o contador) que determina las funciones disponibles en la aplicación.",
    "Rotación de turnos": "Asignación sucesiva y cíclica de los días de riego entre los socios de un pozo según su número de orden.",
    "Servidor": "Equipo o servicio en la nube que procesa las solicitudes de la aplicación y almacena la información del sistema.",
    "Socio": "Ejidatario con derecho al uso del agua de un pozo, que consulta su turno, multas y avisos desde la aplicación.",
    "Turno": "Día asignado a un socio para utilizar el agua del pozo, que inicia y termina a las 18:00 horas.",
    "Usabilidad": "Grado en que un sistema puede ser utilizado por los usuarios con facilidad, eficacia y satisfacción.",
    "Verificador (tester)": "Persona invitada a instalar y probar la aplicación durante la prueba cerrada en Google Play.",
}

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(8)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21.59), Cm(27.94)
s.top_margin = s.bottom_margin = Cm(2.5)
s.left_margin, s.right_margin = Cm(3), Cm(2.5)

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("GLOSARIO DE TÉRMINOS"); r.bold = True; r.font.size = Pt(14)

intro = doc.add_paragraph(
    "A continuación se definen, en orden alfabético, los términos técnicos y propios del contexto comunitario "
    "utilizados a lo largo del presente documento, con el propósito de facilitar su comprensión.")
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

import unicodedata
clave = lambda x: unicodedata.normalize("NFD", x).encode("ascii", "ignore").decode().lower()
letra = ""
for termino in sorted(T, key=clave):
    ini = clave(termino)[0].upper()
    if ini != letra:
        letra = ini
        h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(10); h.paragraph_format.keep_with_next = True
        hr = h.add_run(letra); hr.bold = True; hr.font.size = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(termino + ": ").bold = True
    p.add_run(T[termino])

doc.save("/app/Glosario_de_Terminos.docx")
print("ok", len(T))
