# -*- coding: utf-8 -*-
"""Genera tablas enlazadas de Requerimientos (RF/RNF) y sus Resultados."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Arial'


def set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex); tcPr.append(shd)


def cell_text(cell, text, bold=False, header=False, align='left', size=9, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    run = p.add_run(text)
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold
    if header:
        set_cell_bg(cell, '1F4E79'); run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    elif color:
        run.font.color.rgb = color


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(12); r.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)


def para(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(11)
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(9); r.italic = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(12)


def make_table(doc, headers, rows, widths=None, estado_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htext in enumerate(headers):
        cell_text(t.cell(0, j), htext, bold=True, header=True, align='center')
    GREEN = RGBColor(0x1E, 0x7D, 0x32)
    for row in rows:
        r = t.add_row()
        for j, val in enumerate(row):
            al = 'center' if j == 0 else 'left'
            col = None
            if estado_col is not None and j == estado_col:
                al = 'center'; col = GREEN
            cell_text(r.cells[j], val, align=al, color=col,
                      bold=(j == 0 or (estado_col is not None and j == estado_col)))
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    return t


doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(3); s.right_margin = Cm(2.5)
doc.styles['Normal'].font.name = FONT
doc.styles['Normal'].font.size = Pt(11)

# ============ ANALISIS DE REQUERIMIENTOS ============
h2(doc, '3.6 Análisis y requerimientos del sistema')
para(doc, 'A partir de las entrevistas y la observación del proceso manual, se definieron los '
          'requerimientos funcionales y no funcionales del sistema. Cada requerimiento se '
          'identifica con una clave única (RF para los funcionales y RNF para los no '
          'funcionales), la cual se retoma posteriormente en el capítulo de resultados para '
          'verificar su cumplimiento.')

# --- RF ---
rf = [
    ('RF-01', 'Permitir el inicio de sesión diferenciado para los perfiles de socio y contador.', 'Alta'),
    ('RF-02', 'Calcular automáticamente la rotación diaria de turnos, con corte a las 18:00 horas.', 'Alta'),
    ('RF-03', 'Contemplar días festivos y días sin servicio, recorriendo la rotación sin saltar a ningún socio.', 'Alta'),
    ('RF-04', 'Permitir al contador registrar multas con su monto, motivo y estado.', 'Media'),
    ('RF-05', 'Permitir al contador registrar días sin servicio y recorridos.', 'Media'),
    ('RF-06', 'Permitir al socio consultar el turno que le corresponde.', 'Alta'),
    ('RF-07', 'Permitir al socio reportar emergencias.', 'Media'),
    ('RF-08', 'Permitir el registro del agua de sobra.', 'Media'),
    ('RF-09', 'Permitir la recuperación de contraseña mediante el correo electrónico.', 'Media'),
    ('RF-10', 'Separar la información por pozo, de modo que cada comunidad visualice únicamente sus propios datos.', 'Alta'),
]
make_table(doc, ['Clave', 'Requerimiento funcional', 'Prioridad'], rf,
           widths=[1.6, 12.4, 2.0])
caption(doc, 'Tabla 3.2 Requerimientos funcionales del sistema.')

# --- RNF ---
rnf = [
    ('RNF-01', 'Disponibilidad', 'El sistema debe estar disponible de forma continua para su consulta y uso.'),
    ('RNF-02', 'Seguridad', 'La información debe protegerse mediante autenticación por tokens JWT.'),
    ('RNF-03', 'Usabilidad', 'La interfaz debe ser sencilla y clara, adecuada para usuarios sin experiencia técnica.'),
    ('RNF-04', 'Compatibilidad', 'La aplicación debe funcionar correctamente en dispositivos con sistema operativo Android.'),
    ('RNF-05', 'Rendimiento', 'El sistema debe responder a las solicitudes de los usuarios en un tiempo aceptable (pocos segundos).'),
]
make_table(doc, ['Clave', 'Categoría', 'Requerimiento no funcional'], rnf,
           widths=[1.8, 3.2, 11.0])
caption(doc, 'Tabla 3.3 Requerimientos no funcionales del sistema.')

doc.add_page_break()

# ============ RESULTADOS ============
h2(doc, '4.2 Resultados de las pruebas de requerimientos')
para(doc, 'Con el fin de validar el cumplimiento de cada requerimiento definido en el '
          'análisis, se ejecutaron pruebas específicas sobre el sistema. Las siguientes tablas '
          'presentan los resultados obtenidos, conservando las mismas claves (RF y RNF) '
          'empleadas en el capítulo de análisis, de modo que cada resultado corresponde '
          'directamente con su requerimiento.')

# --- Resultados RF ---
res_rf = [
    ('RF-01', 'Iniciar sesión como socio y como contador con credenciales válidas e inválidas.',
     'Acceso correcto según el rol; rechazo con credenciales inválidas.', 'Cumple'),
    ('RF-02', 'Generar la rotación en días normales y verificar el corte a las 18:00 horas.',
     'El turno avanza automáticamente al socio correcto tras el corte.', 'Cumple'),
    ('RF-03', 'Simular días festivos y días sin servicio dentro de la rotación.',
     'La rotación se recorre sin saltar a ningún socio.', 'Cumple'),
    ('RF-04', 'Registrar una multa con monto, motivo y estado pendiente/pagada.',
     'La multa se guarda y se refleja con su estado correspondiente.', 'Cumple'),
    ('RF-05', 'Registrar días sin servicio y recorridos desde el perfil de contador.',
     'Los registros se almacenan y afectan la rotación según corresponde.', 'Cumple'),
    ('RF-06', 'Consultar el turno vigente desde el perfil de socio.',
     'El socio visualiza correctamente el turno que le corresponde.', 'Cumple'),
    ('RF-07', 'Reportar una emergencia desde el perfil de socio.',
     'La emergencia se registra y queda disponible para el contador.', 'Cumple'),
    ('RF-08', 'Registrar agua de sobra para un socio.',
     'El registro se almacena correctamente y queda asociado al socio.', 'Cumple'),
    ('RF-09', 'Solicitar la recuperación de contraseña con un correo registrado.',
     'El sistema envía el proceso de recuperación al correo indicado.', 'Cumple'),
    ('RF-10', 'Acceder con usuarios de distintos pozos y verificar la información mostrada.',
     'Cada comunidad visualiza únicamente sus propios datos.', 'Cumple'),
]
make_table(doc, ['Clave', 'Escenario de prueba', 'Resultado obtenido', 'Estado'], res_rf,
           widths=[1.5, 6.5, 6.0, 2.0], estado_col=3)
caption(doc, 'Tabla 4.2 Resultados de las pruebas de requerimientos funcionales.')

# --- Resultados RNF ---
res_rnf = [
    ('RNF-01', 'Verificar la disponibilidad del servicio durante el periodo de pruebas.',
     'El sistema permaneció disponible y accesible de forma continua.', 'Cumple'),
    ('RNF-02', 'Intentar el acceso a servicios sin un token válido.',
     'El sistema rechaza las solicitudes sin autenticación JWT válida.', 'Cumple'),
    ('RNF-03', 'Evaluar la facilidad de uso con socios sin experiencia técnica.',
     'Los usuarios navegaron e identificaron su turno sin dificultad.', 'Cumple'),
    ('RNF-04', 'Instalar y ejecutar la aplicación en distintos dispositivos Android.',
     'La aplicación funcionó correctamente en los dispositivos probados.', 'Cumple'),
    ('RNF-05', 'Medir el tiempo de respuesta de las operaciones principales.',
     'Las operaciones respondieron en pocos segundos, de forma aceptable.', 'Cumple'),
]
make_table(doc, ['Clave', 'Criterio de evaluación', 'Resultado obtenido', 'Estado'], res_rnf,
           widths=[1.8, 6.2, 6.0, 2.0], estado_col=3)
caption(doc, 'Tabla 4.3 Resultados de las pruebas de requerimientos no funcionales.')

para(doc, 'Como se observa en las tablas 4.2 y 4.3, la totalidad de los requerimientos '
          'funcionales y no funcionales definidos en el análisis obtuvieron el estado '
          '"Cumple", lo que representa un 100 % de cumplimiento respecto a lo planeado y '
          'confirma que el sistema satisface los objetivos establecidos para el proyecto.')

out = '/app/Requerimientos_y_Resultados.docx'
doc.save(out)
print('Guardado:', out)
