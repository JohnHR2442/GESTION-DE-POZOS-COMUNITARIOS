from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = "2F5597"
GRIS = "D9D9D9"
LINEA = "______________________________"

semanas = [("S", "1", "21-27 sep"), ("S", "2", "28 sep-4 oct"), ("O", "3", "5-11 oct"),
           ("O", "4", "12-18 oct"), ("O", "5", "19-21 oct")]

actividades = [
    ("Levantamiento de requerimientos con el Comisariado Ejidal", [1]),
    ("Análisis y diseño del sistema (casos de uso y modelo de datos)", [1, 2]),
    ("Diseño e implementación de la base de datos", [2]),
    ("Desarrollo del servidor y API REST con autenticación", [2, 3]),
    ("Desarrollo de la aplicación móvil (interfaces y módulos)", [2, 3, 4]),
    ("Implementación del modelo de cálculo de turnos", [3]),
    ("Integración aplicación-servidor y despliegue en la nube", [3, 4]),
    ("Pruebas de funcionamiento y validación", [4]),
    ("Corrección de errores y ajustes finales", [4, 5]),
    ("Generación de versión Android y publicación en prueba cerrada", [5]),
    ("Elaboración de manual de usuario", [4, 5]),
    ("Capacitación de usuario final (Comisariado y socios)", [5]),
    ("Entrega del avance de la documentación y reporte final de residencia", [2, 4, 5]),
]


def sombra(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def texto(celda, t, tam=9, negrita=False, centro=False, cursiva=False, color=None):
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if centro:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    r.font.size = Pt(tam)
    r.font.name = "Arial"
    r.bold = negrita
    r.italic = cursiva
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def alto(fila, cm):
    trPr = fila._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    trPr.append(h)


def anchos(tabla, lista):
    for fila in tabla.rows:
        for i, w in enumerate(lista):
            fila.cells[i].width = Cm(w)


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(27.94), Cm(21.59)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(2.0))
sec.top_margin = sec.bottom_margin = Cm(1.8)


def titulo(t, tam):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(tam)


titulo("FORMATO SEGUIMIENTO DE PROYECTO DE RESIDENCIA PROFESIONAL", 12)
titulo("CRONOGRAMA PRELIMINAR DE ACTIVIDADES", 11)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("(Entregado por el asesor interno)")
r.italic = True
r.font.size = Pt(9)


def dato(etiquetas):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    for i, (et, val) in enumerate(etiquetas):
        if i:
            p.add_run("        ")
        a = p.add_run(et + " ")
        a.bold = True
        a.font.size = Pt(9.5)
        b = p.add_run(val)
        b.bold = True
        b.underline = val != LINEA and val[:1] != "_"
        b.font.size = Pt(9.5)


dato([("ESTUDIANTE:", "Juan José Huerta Rojas"), ("No. DE CONTROL:", "____________________")])
dato([("NOMBRE DEL PROYECTO:", "SISTEMA INTEGRAL DE GESTIÓN DE POZOS COMUNITARIOS")])
dato([("ASESOR(A) EXTERNO(A):", LINEA + "____________")])
dato([("ASESOR(A) INTERNO(A):", LINEA),
      ("PERÍODO DE REALIZACIÓN:", "21 de septiembre al 21 de octubre del 2026")])

# ---------- Tabla principal ----------
n_sem = len(semanas)
filas = 3 + 2 * len(actividades)
t = doc.add_table(rows=filas, cols=2 + n_sem)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
W = [11.5, 1.4] + [2.6] * n_sem
anchos(t, W)

# Encabezado
act = t.cell(0, 0).merge(t.cell(2, 0))
texto(act, "ACTIVIDAD", 11, True, True, cursiva=True)
sombra(act, GRIS)
for f, et in enumerate(["Mes", "Sem", "Fechas"]):
    texto(t.cell(f, 1), et, 8, True, False)
    sombra(t.cell(f, 1), GRIS)
for j, (mes, num, fechas) in enumerate(semanas):
    c = 2 + j
    texto(t.cell(0, c), mes, 9, True, True)
    texto(t.cell(1, c), num, 9, True, True)
    texto(t.cell(2, c), fechas, 7, False, True)
    for f in range(3):
        sombra(t.cell(f, c), GRIS)

# Actividades
for i, (nombre, sems) in enumerate(actividades):
    f = 3 + 2 * i
    c = t.cell(f, 0).merge(t.cell(f + 1, 0))
    texto(c, nombre, 9)
    texto(t.cell(f, 1), "P", 8, centro=True)
    texto(t.cell(f + 1, 1), "R", 8, centro=True)
    for s in sems:
        sombra(t.cell(f, 1 + s), AZUL)
        sombra(t.cell(f + 1, 1 + s), AZUL)
for fila in t.rows:
    alto(fila, 0.42)

doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ---------- Observaciones y entrega de reportes ----------
reportes = ["Reporte 1\n(Semana 2)", "Reporte 2\n(Semana 4)", "Reporte final\n(Semana 5)"]
t2 = doc.add_table(rows=6, cols=2 + len(reportes))
t2.style = "Table Grid"
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
anchos(t2, [3.8, 9.1] + [4.333] * len(reportes))

c = t2.cell(0, 0).merge(t2.cell(0, 1))
texto(c, "", 9)
sombra(c, GRIS)
for j, rep in enumerate(reportes):
    texto(t2.cell(0, 2 + j), rep, 8, True, True)
    sombra(t2.cell(0, 2 + j), GRIS)

c = t2.cell(1, 0).merge(t2.cell(1, 1))
texto(c, "OBSERVACIONES:", 9, True)
alto(t2.rows[1], 1.3)

c = t2.cell(2, 0).merge(t2.cell(5, 0))
texto(c, "ENTREGA\nDE REPORTES", 9, True)
firmantes = ["Docente: " + LINEA, "Estudiante: Juan José Huerta Rojas",
             "Jefe(a) División: " + LINEA, "Asesor(a) Externo(a): " + LINEA]
for k, fi in enumerate(firmantes):
    texto(t2.cell(2 + k, 1), fi, 9)
    alto(t2.rows[2 + k], 0.9)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
r = p.add_run("P = Planeado     R = Realizado     Las celdas sombreadas indican la semana en que se ejecutó la actividad.")
r.font.size = Pt(8)
r.italic = True

doc.save("/app/Cronograma_Residencia.docx")
print("ok")
