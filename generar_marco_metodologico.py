# -*- coding: utf-8 -*-
"""Genera el Capitulo III. Marco Metodologico en Word con formato profesional."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Arial'


def set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, header=False, align='left', size=10):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if header:
        set_cell_bg(cell, '1F4E79')
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_after = Pt(12)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def para(doc, text, justify=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(11)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullet(doc, label, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if label:
        r = p.add_run(label + ' ')
        r.font.name = FONT; r.font.size = Pt(11); r.bold = True
    r2 = p.add_run(text)
    r2.font.name = FONT; r2.font.size = Pt(11)
    return p


doc = Document()
# margenes
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(3); s.right_margin = Cm(2.5)

base = doc.styles['Normal']
base.font.name = FONT
base.font.size = Pt(11)

h1(doc, 'CAPÍTULO III. MARCO METODOLÓGICO')

para(doc, 'En el presente capítulo se describe la metodología empleada para el '
          'desarrollo del proyecto, detallando el tipo y diseño de la investigación, la '
          'población y muestra consideradas, los instrumentos y técnicas de recolección de '
          'datos, las herramientas utilizadas para su procesamiento y, finalmente, los '
          'recursos materiales y humanos que hicieron posible su realización.')

# 3.1
h2(doc, '3.1 Tipo de diseño de investigación')
para(doc, 'La presente investigación se desarrolló bajo un enfoque cuantitativo, debido a '
          'que la recolección y el análisis de la información se fundamentan en datos medibles '
          'y en el registro objetivo del comportamiento del sistema y de sus usuarios. El '
          'diseño es de tipo no experimental, ya que las variables se observaron en su '
          'contexto natural sin ser manipuladas de forma deliberada, y transversal, puesto que '
          'los datos se recolectaron en un periodo único y determinado.')
para(doc, 'Por su alcance, el estudio es descriptivo y correlacional. Es descriptivo porque '
          'caracteriza de manera detallada el proceso de administración de los turnos de agua '
          'en las comunidades de San Isidro, Zapata y Cárdenas, identificando sus '
          'particularidades, actores y deficiencias. Es correlacional porque pretende '
          'establecer la relación existente entre la implementación de la aplicación móvil '
          '(variable independiente) y el nivel de orden y transparencia alcanzado en la '
          'distribución del agua (variable dependiente).')
para(doc, 'Para la construcción de la solución tecnológica se adoptó la metodología de '
          'desarrollo incremental, la cual permite dividir el proyecto en incrementos '
          'funcionales que se diseñan, programan y prueban de manera sucesiva. Esta elección '
          'resultó conveniente porque posibilitó entregar versiones operativas tempranas, '
          'recibir retroalimentación de los contadores y socios, y ajustar los requerimientos '
          'de forma progresiva hasta obtener el producto final publicado. Las fases '
          'contempladas fueron: análisis de requerimientos, diseño de la arquitectura y la '
          'base de datos, codificación, pruebas e integración, y despliegue en producción.')

# 3.2
h2(doc, '3.2 Población y muestra')
para(doc, 'La población de estudio estuvo conformada por la totalidad de los socios y '
          'contadores asociados a los tres pozos comunitarios administrados por el proyecto. '
          'En conjunto, dicha población asciende a aproximadamente sesenta socios activos y '
          'tres contadores, correspondiendo un contador a cada pozo (San Isidro, Zapata y '
          'Cárdenas). Esta población representa al total de personas directamente involucradas '
          'en el proceso de distribución del agua y, por lo tanto, en el uso del sistema '
          'desarrollado.')
para(doc, 'Debido a que el universo de estudio es reducido, finito y plenamente accesible, '
          'se optó por un muestreo no probabilístico por conveniencia. Bajo este criterio, la '
          'muestra quedó integrada por aquellos socios y contadores que aceptaron participar '
          'de manera voluntaria en la fase de pruebas de la aplicación, instalándola en sus '
          'dispositivos y utilizándola en condiciones reales de operación. Este tipo de '
          'muestreo es pertinente en investigaciones de carácter tecnológico aplicado, donde '
          'el objetivo es validar el funcionamiento del sistema con los usuarios reales a '
          'quienes está destinado, más que generalizar estadísticamente a una población '
          'amplia.')

# tabla poblacion
t = doc.add_table(rows=5, cols=3)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_text(t.cell(0, 0), 'Pozo comunitario', bold=True, header=True, align='center')
cell_text(t.cell(0, 1), 'Socios (aprox.)', bold=True, header=True, align='center')
cell_text(t.cell(0, 2), 'Contadores', bold=True, header=True, align='center')
datos = [('San Isidro', '25', '1'), ('Zapata', '20', '1'), ('Cárdenas', '15', '1'),
         ('Total', '60', '3')]
for i, (a, b, c) in enumerate(datos, start=1):
    bold = (a == 'Total')
    cell_text(t.cell(i, 0), a, bold=bold, align='left')
    cell_text(t.cell(i, 1), b, bold=bold, align='center')
    cell_text(t.cell(i, 2), c, bold=bold, align='center')
cap = doc.add_paragraph()
cr = cap.add_run('Tabla 3.1 Distribución de la población por pozo comunitario.')
cr.font.name = FONT; cr.font.size = Pt(9); cr.italic = True
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after = Pt(10)

# 3.3
h2(doc, '3.3 Instrumentos, técnicas y métodos')
para(doc, 'Con la finalidad de obtener la información necesaria para el análisis y la '
          'validación del sistema, se emplearon diversas técnicas e instrumentos de '
          'recolección de datos, descritos a continuación:')
bullet(doc, 'Entrevista:', 'se aplicaron entrevistas semiestructuradas a los tres contadores '
       'de los pozos, con el propósito de conocer a detalle las reglas de rotación de turnos, '
       'los horarios de corte, el manejo de días festivos y sin servicio, así como el '
       'procedimiento para la asignación y cobro de multas.')
bullet(doc, 'Observación directa:', 'se observó en campo la manera en que se llevaba a cabo '
       'la administración manual de los turnos, lo que permitió identificar las principales '
       'fuentes de error, duplicidad de información y conflictos entre socios.')
bullet(doc, 'Análisis documental:', 'se revisaron las libretas, listas y registros físicos '
       'utilizados por los contadores, con el fin de comprender la lógica de la rotación y '
       'trasladarla correctamente al algoritmo del sistema.')
bullet(doc, 'Cuestionario:', 'se diseñó y aplicó un cuestionario basado en una escala de '
       'Likert a los usuarios participantes, para medir el nivel de aceptación, usabilidad y '
       'satisfacción respecto a la aplicación.')
para(doc, 'En cuanto a los métodos, se recurrió al método analítico para descomponer el '
          'problema general en requerimientos específicos del sistema, y al método de '
          'desarrollo incremental de software, que estructuró el trabajo técnico en fases '
          'claramente definidas de análisis, diseño, programación, pruebas y despliegue. La '
          'combinación de ambos métodos permitió abordar tanto la dimensión investigativa '
          'como la dimensión de ingeniería del proyecto.')

# 3.4
h2(doc, '3.4 Procesamiento de datos')
para(doc, 'El procesamiento de la información se apoyó en un conjunto de herramientas y '
          'tecnologías de software, cada una con una función específica dentro de la '
          'arquitectura y el flujo de datos del sistema:')
bullet(doc, 'MongoDB Atlas:', 'base de datos no relacional alojada en la nube, encargada de '
       'almacenar y organizar la información de pozos, usuarios, turnos, días sin servicio, '
       'emergencias y multas.')
bullet(doc, 'FastAPI (Python):', 'framework del lado del servidor que recibe las solicitudes, '
       'ejecuta la lógica de negocio —incluido el cálculo automático de turnos— y expone los '
       'servicios mediante una API REST segura con autenticación por tokens JWT.')
bullet(doc, 'React Native y Expo:', 'tecnologías del cliente móvil responsables de presentar '
       'la información procesada y de capturar las acciones del usuario de forma intuitiva.')
bullet(doc, 'Postman:', 'herramienta utilizada para probar, depurar y validar el correcto '
       'funcionamiento de cada uno de los servicios de la API REST.')
bullet(doc, 'Microsoft Excel:', 'empleado para tabular, calcular frecuencias y porcentajes, '
       'y graficar los resultados obtenidos de los cuestionarios de satisfacción.')
bullet(doc, 'Git y GitHub:', 'sistema de control de versiones utilizado para el respaldo, '
       'seguimiento de cambios y gestión del código fuente a lo largo del desarrollo.')
para(doc, 'El procesamiento de los datos operativos generados por los usuarios (turnos, '
          'multas y agua de sobra) se realiza de manera automática dentro del propio sistema, '
          'garantizando la integridad y disponibilidad de la información. Por su parte, los '
          'datos provenientes de los cuestionarios se analizaron mediante estadística '
          'descriptiva, obteniendo frecuencias y porcentajes que permitieron interpretar el '
          'nivel de aceptación de la solución.')

# 3.5
h2(doc, '3.5 Recursos materiales y de capital humano')
para(doc, 'Para la ejecución del proyecto se requirieron recursos materiales —hardware, '
          'software y servicios— así como recursos humanos, los cuales se detallan en las '
          'siguientes tablas.')

# Tabla recursos materiales
h3 = doc.add_paragraph()
r = h3.add_run('Recursos materiales')
r.font.name = FONT; r.font.size = Pt(11); r.bold = True
h3.paragraph_format.space_before = Pt(6); h3.paragraph_format.space_after = Pt(4)

tm = doc.add_table(rows=1, cols=2)
tm.style = 'Table Grid'
tm.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_text(tm.cell(0, 0), 'Recurso', bold=True, header=True, align='center')
cell_text(tm.cell(0, 1), 'Descripción / Uso', bold=True, header=True, align='center')
mat = [
    ('Computadora portátil', 'Equipo de desarrollo, programación y pruebas del sistema.'),
    ('Teléfonos Android', 'Dispositivos para la instalación y prueba de la aplicación en campo.'),
    ('Conexión a internet', 'Enlace permanente para el desarrollo, despliegue y uso del sistema.'),
    ('Visual Studio Code', 'Entorno de desarrollo integrado para la escritura del código.'),
    ('Node.js / Python', 'Entornos de ejecución del cliente (Expo) y del servidor (FastAPI).'),
    ('Expo y EAS CLI', 'Herramientas para la compilación y generación de la aplicación Android.'),
    ('MongoDB Atlas', 'Servicio de base de datos en la nube (plan gratuito).'),
    ('Render', 'Servicio de hospedaje en la nube para el servidor backend.'),
    ('Google Play Console', 'Plataforma para la publicación y distribución de la aplicación.'),
]
for a, b in mat:
    row = tm.add_row()
    cell_text(row.cells[0], a, bold=True, align='left')
    cell_text(row.cells[1], b, align='left')

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# Tabla recursos humanos
h4 = doc.add_paragraph()
r = h4.add_run('Recursos de capital humano')
r.font.name = FONT; r.font.size = Pt(11); r.bold = True
h4.paragraph_format.space_before = Pt(6); h4.paragraph_format.space_after = Pt(4)

th = doc.add_table(rows=1, cols=2)
th.style = 'Table Grid'
th.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_text(th.cell(0, 0), 'Participante', bold=True, header=True, align='center')
cell_text(th.cell(0, 1), 'Función en el proyecto', bold=True, header=True, align='center')
hum = [
    ('Residente (desarrollador)', 'Análisis, diseño, programación, pruebas y publicación de la aplicación.'),
    ('Asesor interno (docente)', 'Supervisión académica y metodológica del proyecto.'),
    ('Contadores de los pozos', 'Aportación de las reglas de distribución y validación del sistema.'),
    ('Socios de las comunidades', 'Participación en las pruebas de usabilidad y funcionamiento.'),
]
for a, b in hum:
    row = th.add_row()
    cell_text(row.cells[0], a, bold=True, align='left')
    cell_text(row.cells[1], b, align='left')

out = '/app/Capitulo_III_Marco_Metodologico.docx'
doc.save(out)
print('Guardado:', out)
