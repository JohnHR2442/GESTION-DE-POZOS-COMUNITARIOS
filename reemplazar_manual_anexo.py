"""Reemplaza el manual del Anexo A de la tesis por la versión nueva.
Uso: python reemplazar_manual_anexo.py tesis.docx manual.docx salida.docx [offset_paginas]"""
import sys, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docxcompose.composer import Composer

TESIS, MANUAL, OUT = sys.argv[1:4]
OFFSET = int(sys.argv[4]) if len(sys.argv) > 4 else 0

# ---------- 1. Preparar el manual (estilos propios para no mezclarse con los índices de la tesis)
m = Document(MANUAL)
for viejo, nuevo, tam in (("Heading 1", "Manual Titulo 1", 14), ("Heading 2", "Manual Titulo 2", 12)):
    s = m.styles.add_style(nuevo, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = m.styles["Normal"]
    s.font.name = "Arial"; s.font.size = Pt(tam); s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.keep_with_next = True
    for p in m.paragraphs:
        if p.style.name == viejo:
            p.style = s
for viejo, nuevo in (("Cap_Figura", "Manual Figura"), ("Cap_Tabla", "Manual Tabla")):
    s = m.styles[viejo]
    s.name = nuevo
    s.element.set(qn("w:styleId"), nuevo.replace(" ", ""))
    s.font.name = "Arial"; s.font.size = Pt(9); s.font.italic = False
    for p in m.element.body.iter(qn("w:pStyle")):
        if p.get(qn("w:val")) == viejo:
            p.set(qn("w:val"), nuevo.replace(" ", ""))
for it in m.element.body.iter(qn("w:instrText")):
    t = it.text or ""
    t = t.replace('\\o "1-2"', '\\t "Manual Titulo 1,1,Manual Titulo 2,2"')
    t = t.replace('"Cap_Figura,1"', '"Manual Figura,1"').replace('"Cap_Tabla,1"', '"Manual Tabla,1"')
    t = re.sub(r"_Toc(\d+)", r"_TocM\1", t)
    it.text = t
for el in m.element.body.iter():
    for atr in (qn("w:name"), qn("w:anchor")):
        v = el.get(atr)
        if v and v.startswith("_Toc"):
            el.set(atr, "_TocM" + v[4:])
# Tipo de letra de la tesis en todo el manual: Arial
for r in m.element.body.iter(qn("w:r")):
    rPr = r.find(qn("w:rPr"))
    if rPr is None:
        rPr = r.makeelement(qn("w:rPr"), {}); r.insert(0, rPr)
    f = rPr.find(qn("w:rFonts"))
    if f is None:
        f = rPr.makeelement(qn("w:rFonts"), {}); rPr.insert(0, f)
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        f.set(qn("w:" + k), "Arial")
# Número de página en los índices internos del manual (lugar dentro de la tesis)
for p in m.paragraphs:
    if p.style.name.startswith("toc") or p.style.name == "Indice_Figuras_Tablas":
        ts = [x for x in p._p.iter(qn("w:t")) if x.text and x.text.strip()]
        if ts and ts[-1].text.strip().isdigit():
            ts[-1].text = str(int(ts[-1].text.strip()) + OFFSET)
# Sin encabezado/pie propios: se usan los de la tesis
sect = m.element.body.find(qn("w:sectPr"))
for ref in sect.findall(qn("w:headerReference")) + sect.findall(qn("w:footerReference")):
    sect.remove(ref)
tp = sect.find(qn("w:titlePg"))
if tp is not None:
    sect.remove(tp)
m.save("/tmp/manual_prep2.docx")

# ---------- 2. Quitar el manual anterior del Anexo A
t = Document(TESIS)
body = t.element.body
kids = list(body)
txt = lambda e: "".join(x.text or "" for x in e.iter(qn("w:t"))).strip()
i_anexo = max(i for i, e in enumerate(kids) if e.tag == qn("w:p") and txt(e).startswith("ANEXO A."))
i_salto = next(i for i in range(i_anexo, len(kids)) if 'w:type="page"' in kids[i].xml)
i_fin = next(i for i in range(i_salto + 1, len(kids)) if txt(kids[i]).startswith("Formato de títulos"))
assert 'w:type="page"' in kids[i_fin].xml
viejos = kids[i_salto + 1:i_fin]
assert any("MANUAL DE USUARIO" in txt(e) for e in viejos[:6]) and "Soporte y contacto" in "".join(txt(e) for e in viejos[-8:])
for e in viejos:
    body.remove(e)
print("Eliminados", len(viejos), "elementos del manual anterior")

# ---------- 3. Insertar el manual nuevo en su lugar
comp = Composer(t)
comp.insert(i_salto + 1, Document("/tmp/manual_prep2.docx"))
comp.save(OUT)
print("OK", OUT)
