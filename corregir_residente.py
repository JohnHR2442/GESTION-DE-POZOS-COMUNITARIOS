import io
import zipfile
from PIL import Image
from docx import Document

SRC = "/app/tesis_v3.docx"
TMP = "/tmp/tesis_v3_txt.docx"
OUT = "/app/Tesis_Corregida.docx"

d = Document(SRC)
P = d.paragraphs

cambios = [
    (153, "El estudiante se integró formalmente",
     "El autor, egresado de Ingeniería en Sistemas Computacionales, se integró formalmente"),
    (182, "el Estudiante Residente, en su función",
     "el Desarrollador del Sistema (autor del proyecto), en su función"),
    (205, "dentro del periodo establecido para la residencia profesional",
     "en el periodo del 21 de septiembre al 21 de octubre de 2026"),
]


def reemplazar(par, viejo, nuevo):
    for r in par.runs:
        if viejo in r.text:
            r.text = r.text.replace(viejo, nuevo)
            return True
    # texto repartido en varios runs: se concentra en el primero
    if viejo in par.text:
        t = par.text.replace(viejo, nuevo)
        par.runs[0].text = t
        for r in par.runs[1:]:
            r.text = ""
        return True
    return False


for i, viejo, nuevo in cambios:
    assert reemplazar(P[i], viejo, nuevo), (i, viejo)

# Conclusiones: relacionar el cumplimiento del objetivo con la hipótesis
assert P[599].text.startswith("Sobre el cumplimiento del objetivo")
ult = P[599].runs[-1]
ult.text = ult.text.rstrip() + (" De acuerdo con estos resultados, se acepta la hipótesis de investigación (Hi) y se "
                                "rechaza la hipótesis nula (H0), ya que la aplicación mejoró el orden y la "
                                "transparencia en la distribución del agua de los pozos San Isidro, Zapata y Cárdenas.")

# Tablas
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                reemplazar(par, "Residente (desarrollador)", "Desarrollador (autor del proyecto)")
                reemplazar(par, "reporte final de residencia", "reporte final del proyecto")

d.save(TMP)

# Reemplazar la imagen del organigrama (word/media/image2.jpeg) conservando el mismo tamaño
img = Image.open("/app/organigrama_comisariado.png").convert("RGB").resize((1568, 1176), Image.LANCZOS)
buf = io.BytesIO()
img.save(buf, "JPEG", quality=95)
with zipfile.ZipFile(TMP) as zin, zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = buf.getvalue() if item.filename == "word/media/image2.jpeg" else zin.read(item.filename)
        zout.writestr(item, data)
print("ok")
