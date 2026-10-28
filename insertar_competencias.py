"""Inserta el apartado de Competencias en la tesis respetando el formato del documento.
Uso: python insertar_competencias.py tesis.docx competencias.docx salida.docx"""
import sys, copy, io, re
from docx import Document
from docx.shared import Pt, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

TESIS, COMP, OUT = sys.argv[1:4]
t = Document(TESIS)
c = Document(COMP)
P = t.paragraphs

# --- Ubicar el apartado en la tesis
h2 = next(p for p in P if p.style.name == 'Heading 2' and 'COMPETENCIAS DESARROLLADAS' in p.text.upper())
placeholders = []
el = h2._p.getnext()
while el is not None and el.tag == qn('w:p') and Paragraph(el, h2._parent).style.name == 'Subtitle':
    placeholders.append(el)
    el = el.getnext()
modelo_sub = next(e for e in placeholders if ''.join(e.itertext()).strip())
modelo_capt = next(p for p in P if p.style.name == 'Cap_Tabla' and p.text.startswith('Tabla 12'))._p
modelo_capf = next(p for p in P if p.style.name == 'Cap_Figura' and p.text.startswith('Figura 13'))._p
ancla = placeholders[0]


def insertar(elem):
    ancla.addprevious(elem)


def arial(r, size=None, bold=None):
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    r.font.size = Pt(size) if size else None
    if bold is not None:
        r.bold = bold


def nuevo_p(estilo='Normal'):
    p = Paragraph(OxmlElement('w:p'), h2._parent)
    p.style = t.styles[estilo]
    return p


def subtitulo(texto):
    e = copy.deepcopy(modelo_sub)
    for r in e.findall(qn('w:r'))[1:]:
        e.remove(r)
    for r in e.findall(qn('w:r')):
        e.remove(r)
    p = Paragraph(e, h2._parent)
    r = p.add_run(texto.rstrip('.') + '.')
    arial(r, 12, True)
    r.italic = False
    r.font.color.rgb = RGBColor(0, 0, 0)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '0'); r._element.rPr.append(sp)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    insertar(e)


def subtema(texto):
    p = nuevo_p()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.5
    pf.space_after = Pt(12)
    pf.first_line_indent = None
    arial(p.add_run(texto), 11, True)
    insertar(p._p)


def parrafo(src, estilo='Normal'):
    p = nuevo_p(estilo)
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_after = Pt(12 if estilo == 'Normal' else 6)
    for r in src.runs:
        arial(p.add_run(r.text), 11, bool(r.bold))
    insertar(p._p)


BM = [900100]


def pie(modelo, texto):
    e = copy.deepcopy(modelo)
    for ch in list(e):
        if ch.tag != qn('w:pPr'):
            e.remove(ch)
    p = Paragraph(e, h2._parent)
    BM[0] += 1
    nombre = '_TocComp%d' % BM[0]
    ini = OxmlElement('w:bookmarkStart'); ini.set(qn('w:id'), str(BM[0])); ini.set(qn('w:name'), nombre)
    fin = OxmlElement('w:bookmarkEnd'); fin.set(qn('w:id'), str(BM[0]))
    e.append(ini)
    r = p.add_run(texto)
    arial(r, 9)
    r.italic = False
    e.append(fin)
    MARCAS[texto] = nombre
    insertar(e)


MARCAS = {}


def fuente():
    p = nuevo_p()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    arial(p.add_run('Fuente: Elaboración propia.'), 9)
    insertar(p._p)


def tabla(tbl_el):
    e = copy.deepcopy(tbl_el)
    for rpr in e.iter(qn('w:rPr')):
        f = rpr.find(qn('w:rFonts'))
        if f is None:
            f = OxmlElement('w:rFonts'); rpr.insert(0, f)
        for k in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            f.set(qn(f'w:{k}'), 'Arial')
    for r in e.iter(qn('w:r')):
        if r.find(qn('w:rPr')) is None:
            rpr = OxmlElement('w:rPr'); f = OxmlElement('w:rFonts')
            for k in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
                f.set(qn(f'w:{k}'), 'Arial')
            rpr.append(f); r.insert(0, rpr)
    insertar(e)


def imagen(src_p):
    blip = src_p._p.find('.//' + qn('a:blip'))
    blob = c.part.related_parts[blip.get(qn('r:embed'))].blob
    ext = src_p._p.find('.//' + qn('wp:extent'))
    p = nuevo_p()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(io.BytesIO(blob), width=Emu(int(ext.get('cx'))))
    insertar(p._p)


# --- Recorrer el documento de competencias (se omite el título, la tesis ya lo tiene)
SUBTITULOS = ('Competencias desarrolladas', 'Aporte al perfil de egreso', 'Grafos del proyecto')
pendiente = None  # pie de tabla que va debajo de la tabla
elementos = list(c.element.body.iterchildren())
for e in elementos[1:]:
    if e.tag == qn('w:tbl'):
        tabla(e)
        if pendiente:
            pie(modelo_capt, pendiente); pendiente = None
        continue
    if e.tag != qn('w:p'):
        continue
    p = Paragraph(e, c)
    txt = p.text.strip()
    if 'graphic' in e.xml:
        imagen(p)
    elif not txt or txt.startswith('Fuente:'):
        if txt:
            fuente()
    elif txt in SUBTITULOS:
        subtitulo(txt)
    elif txt.startswith('Competencias ') and len(txt) < 40:
        subtema(txt)
    elif re.match(r'Tabla \d+\.', txt):
        pendiente = txt
    elif re.match(r'Figura \d+\.', txt):
        pie(modelo_capf, txt)
    else:
        parrafo(p, 'List Bullet' if p.style.name == 'List Bullet' else 'Normal')

for e in placeholders:
    e.getparent().remove(e)

# --- Agregar entradas a los índices de figuras y tablas
def agregar_indice(prefijo_ultimo, nuevos):
    ult = next(p for p in t.paragraphs if p.style.name.startswith('toc') and p.text.startswith(prefijo_ultimo))._p
    for texto in nuevos:
        e = copy.deepcopy(ult)
        ts = list(e.iter(qn('w:t')))
        ts[0].text = texto
        ts[-1].text = '0'
        e.find('.//' + qn('w:hyperlink')).set(qn('w:anchor'), MARCAS[texto])
        e.find('.//' + qn('w:instrText')).text = ' PAGEREF %s \\h ' % MARCAS[texto]
        ult.addnext(e)
        ult = e

caps = [x.strip() for x in (pp.text for pp in c.paragraphs)]
agregar_indice('Figura 13.', [x for x in caps if re.match(r'Figura \d+\.', x)])
agregar_indice('Tabla 14.', [x for x in caps if re.match(r'Tabla \d+\.', x)])

t.save(OUT)
print('OK', OUT)
