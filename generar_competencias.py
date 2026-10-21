from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = "1F3A5F"
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(8)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3), Cm(2.5)


def sombra(c, color):
    tcPr = c._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
    tcPr.append(s)


def titulo(t, tam=12, centro=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(t); r.bold = True; r.font.size = Pt(tam)


def p(t):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.add_run(t)
    return par


def vineta(t, neg=None):
    par = doc.add_paragraph(style="List Bullet")
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(3)
    if neg:
        par.add_run(neg).bold = True
    par.add_run(t)


def caption(t):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_after = Pt(4)
    r = par.add_run(t); r.font.size = Pt(9)


def fuente():
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run("Fuente: Elaboración propia.")
    r.font.size = Pt(9)


def tabla(enc, filas, anchos):
    t = doc.add_table(rows=1 + len(filas), cols=len(enc))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, e in enumerate(enc):
        c = t.cell(0, j); sombra(c, AZUL)
        pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.line_spacing = 1.0
        r = pp.add_run(e); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(255, 255, 255)
    for i, f in enumerate(filas, 1):
        for j, v in enumerate(f):
            c = t.cell(i, j)
            pp = c.paragraphs[0]
            pp.paragraph_format.line_spacing = 1.0
            pp.paragraph_format.space_after = Pt(2)
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = pp.add_run(v); r.font.size = Pt(9.5)
            if j == 0: r.bold = True
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for fila in t.rows:
        fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, w in enumerate(anchos):
            fila.cells[j].width = Cm(w)
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(th)
    return t


# ================== COMPETENCIAS ==================
titulo("COMPETENCIAS DESARROLLADAS Y/O APLICADAS", 12, centro=True)

titulo("Competencias desarrolladas")
p("Durante el desarrollo del proyecto “Sistema Integral de Gestión de Pozos Comunitarios”, realizado en el "
  "Comisariado Ejidal de San José Carpinteros, se aplicaron y fortalecieron competencias adquiridas a lo largo "
  "de la carrera de Ingeniería en Sistemas Computacionales. Estas competencias permitieron analizar la "
  "problemática de la asignación manual de turnos de agua, diseñar una solución tecnológica y llevarla a "
  "producción mediante una aplicación móvil conectada a un servidor en la nube.")
p("Las competencias se clasifican en específicas, relacionadas directamente con las asignaturas del plan de "
  "estudios, y genéricas, que corresponden a habilidades transversales del perfil profesional.")

titulo("Competencias específicas", 11)
p("En la Tabla 15 se describen las asignaturas cuyas competencias se aplicaron de manera directa en el proyecto "
  "y la forma concreta en que se utilizaron.")
caption("Tabla 15. Competencias específicas aplicadas en el proyecto.")
tabla(["Asignatura", "Competencia aplicada", "Aplicación en el proyecto"], [
    ("Programación Web", "Diseñar y construir servicios web basados en protocolos estándar.",
     "Desarrollo del servidor con FastAPI e implementación de servicios REST (HTTP/JSON) para turnos, multas, emergencias y notificaciones."),
    ("Tópicos Avanzados de Programación", "Desarrollar interfaces gráficas y aplicaciones para dispositivos móviles.",
     "Desarrollo de la aplicación Android con React Native y Expo: pantallas, navegación por pestañas, calendario y modo oscuro."),
    ("Ingeniería de Software", "Aplicar metodologías y modelos para el desarrollo de software de calidad.",
     "Levantamiento de requerimientos funcionales y no funcionales, casos de uso y pruebas de funcionamiento."),
    ("Taller de Base de Datos", "Implementar bases de datos y realizar operaciones de consulta y manipulación.",
     "Implementación de las colecciones de pozos, usuarios, multas, emergencias y días sin servicio en MongoDB."),
    ("Administración de Bases de Datos", "Administrar, respaldar y asegurar la disponibilidad de la información.",
     "Configuración de MongoDB Atlas en la nube, control de acceso y respaldo de la información."),
    ("Fundamentos de Base de Datos", "Modelar los datos de un sistema de información.",
     "Diseño del modelo de datos y de las relaciones entre socios, pozos y registros."),
    ("Matemáticas Discretas", "Aplicar estructuras y operaciones discretas para resolver problemas.",
     "Modelo de aritmética modular para calcular el turno diario y recorrer festivos y días sin servicio."),
    ("Fundamentos de Ingeniería de Software", "Identificar el ciclo de vida del software y sus etapas.",
     "Planeación del proyecto por etapas: análisis, diseño, desarrollo, pruebas e implementación."),
    ("Gestión de Proyectos de Software", "Planificar, organizar y controlar proyectos de software.",
     "Elaboración del cronograma de actividades, seguimiento semanal y publicación en Google Play."),
    ("Redes de Computadoras", "Comprender la comunicación cliente-servidor en redes.",
     "Comunicación de la aplicación con el servidor desplegado en Render mediante internet."),
    ("Seguridad en Redes (especialidad)", "Aplicar mecanismos de seguridad de la información.",
     "Autenticación con JWT, cifrado de contraseñas, roles de usuario y comunicación por HTTPS."),
], [4.0, 5.5, 6.5])
fuente()

titulo("Competencias genéricas", 11)
vineta(" capacidad de análisis y síntesis para identificar la problemática del reparto de agua; "
       "solución de problemas; toma de decisiones técnicas; comunicación oral y escrita en la elaboración "
       "de documentación y del manual de usuario.", "Instrumentales:")
vineta(" trabajo en equipo con el Comisariado Ejidal, socios y contadores; capacidad crítica y autocrítica; "
       "compromiso ético con el manejo de la información de los usuarios.", "Interpersonales:")
vineta(" aplicación de los conocimientos en la práctica; habilidad de investigación; capacidad de aprender "
       "nuevas tecnologías (Expo, FastAPI, MongoDB Atlas, Render y Google Play Console); iniciativa y "
       "búsqueda del logro al llevar el sistema a producción.", "Sistémicas:")

# ================== APORTE AL PERFIL DE EGRESO ==================
titulo("Aporte al perfil de egreso")
p("El proyecto contribuyó al logro de los atributos del perfil de egreso del Ingeniero en Sistemas "
  "Computacionales (plan ISIC-2010-224) establecidos por el Tecnológico Nacional de México, cuyo objetivo "
  "general es formar profesionistas capaces de diseñar, desarrollar, implementar y administrar tecnología "
  "computacional para aportar soluciones innovadoras en beneficio de la sociedad. En la Tabla 16 se relaciona "
  "cada atributo con la evidencia obtenida durante el desarrollo del proyecto.")
caption("Tabla 16. Aporte del proyecto al perfil de egreso.")
tabla(["Atributo del perfil de egreso", "Aporte del proyecto"], [
    ("1. Implementa aplicaciones computacionales para solucionar problemas de diversos contextos, integrando "
     "diferentes tecnologías, plataformas o dispositivos.",
     "Se integró una aplicación móvil Android, un servidor en la nube y una base de datos para resolver la "
     "asignación de turnos de agua de tres pozos comunitarios."),
    ("2. Diseña, desarrolla y aplica modelos computacionales para solucionar problemas, mediante la selección "
     "y uso de herramientas matemáticas.",
     "Se diseñó el modelo de aritmética modular que calcula el turno diario considerando el corte de las 18:00 "
     "horas, los festivos y los días sin servicio."),
    ("3. Diseña e implementa interfaces para la automatización de sistemas de hardware y desarrollo del "
     "software asociado.",
     "Aporte parcial: el proyecto se centró en software; se desarrollaron las interfaces de usuario y la "
     "comunicación con el servidor que automatizan el control de turnos, sin integración de hardware."),
    ("4. Coordina y participa en equipos multidisciplinarios para la aplicación de soluciones innovadoras en "
     "diferentes contextos.",
     "Se trabajó con el Comisariado Ejidal, el Subcomité Técnico de Aguas, contadores y socios para validar "
     "los requerimientos y realizar las pruebas."),
    ("5. Diseña, implementa y administra bases de datos optimizando los recursos disponibles, conforme a las "
     "normas vigentes de manejo y seguridad de la información.",
     "Se diseñó, implementó y administró la base de datos en MongoDB Atlas, con control de acceso por roles y "
     "contraseñas cifradas."),
    ("6. Desarrolla y administra software para apoyar la productividad y competitividad de las organizaciones "
     "cumpliendo con estándares de calidad.",
     "El sistema sustituyó el control manual en libretas, redujo los tiempos de consulta y dio transparencia al "
     "registro de multas del Comisariado."),
    ("7. Evalúa tecnologías de hardware para soportar aplicaciones de manera efectiva.",
     "Se definieron los requisitos mínimos de los teléfonos Android de los usuarios y se evaluaron los recursos "
     "del servidor en la nube (Render) y de la base de datos (MongoDB Atlas) para soportar la aplicación."),
    ("8. Detecta áreas de oportunidad empleando una visión empresarial para crear proyectos aplicando las "
     "Tecnologías de la Información y Comunicación.",
     "Se identificó la falta de control en el reparto de agua como área de oportunidad y se resolvió con una "
     "solución digital de bajo costo."),
    ("9. Diseña, configura y administra redes de computadoras para crear soluciones de conectividad en la "
     "organización, aplicando las normas y estándares vigentes.",
     "Se configuró la comunicación cliente-servidor por internet mediante HTTPS y autenticación con JWT, "
     "permitiendo que los socios consulten la información desde cualquier lugar."),
], [7.5, 8.5])
fuente()

# ================== GRAFO (hoja horizontal) ==================
ns = doc.add_section(WD_SECTION.NEW_PAGE)
ns.orientation = WD_ORIENT.LANDSCAPE
ns.page_width, ns.page_height = Cm(27.94), Cm(21.59)
ns.top_margin = ns.bottom_margin = Cm(2)
ns.left_margin = ns.right_margin = Cm(2.5)

titulo("Grafos del proyecto")
p("La Figura 14 muestra el grafo de la retícula de Ingeniería en Sistemas Computacionales, en el que se "
  "identifica con colores la relación de cada asignatura con el proyecto: las seleccionadas se aplicaron "
  "directamente; las de aporte directo e indirecto apoyaron el desarrollo; las de competencia previa "
  "constituyeron la base de conocimientos, y las de especialidad corresponden al módulo cursado.")
pi = doc.add_paragraph(); pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
pi.paragraph_format.space_after = Pt(2); pi.paragraph_format.line_spacing = 1.0
pi.paragraph_format.keep_with_next = True
pi.add_run().add_picture("/app/grafo_competencias.png", width=Cm(21.5))
caption("Figura 14. Grafo de asignaturas de Ingeniería en Sistemas Computacionales aplicadas en el proyecto.")
fuente()

ns2 = doc.add_section(WD_SECTION.NEW_PAGE)
ns2.orientation = WD_ORIENT.PORTRAIT
ns2.page_width, ns2.page_height = Cm(21.59), Cm(27.94)
ns2.top_margin = ns2.bottom_margin = Cm(2.5)
ns2.left_margin, ns2.right_margin = Cm(3), Cm(2.5)

p("En la Tabla 17 se resume la clasificación de las asignaturas representadas en el grafo.")
caption("Tabla 17. Clasificación de asignaturas según su relación con el proyecto.")
tabla(["Clasificación", "Asignaturas", "Total"], [
    ("Seleccionada", "Programación Web, Tópicos Avanzados de Programación, Ingeniería de Software "
     "y Taller de Base de Datos.", "4"),
    ("Aporte directo", "Matemáticas Discretas, Fundamentos de Base de Datos, Fundamentos de Ingeniería de Software "
     "Redes de Computadoras, Administración de Bases de Datos y Gestión de Proyectos de Software.", "6"),
    ("Aporte indirecto", "Probabilidad y Estadística, Sistemas Operativos, Taller de Investigación I, "
     "Taller de Investigación II y Administración de Redes.", "5"),
    ("Competencia previa directa", "Fundamentos de Programación, Programación Orientada a Objetos y "
     "Estructura de Datos.", "3"),
    ("Competencia previa indirecta", "Fundamentos de Investigación, Taller de Ética, Taller de Administración, "
     "Contabilidad Financiera, Cultura Empresarial y Desarrollo Sustentable.", "6"),
    ("Especialidad", "Seguridad en Redes, Programación de Robots, Redes Neuronales Artificiales y "
     "Arquitecturas Reconfigurables.", "4"),
], [4.0, 10.5, 1.5])
fuente()
p("De las asignaturas de especialidad, Seguridad en Redes fue la de mayor aplicación, al utilizarse en la "
  "autenticación de usuarios y en la protección de la información del sistema.")

doc.save("/app/Competencias_Desarrolladas.docx")
print("ok")
