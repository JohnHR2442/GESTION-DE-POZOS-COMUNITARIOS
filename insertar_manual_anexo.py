import copy
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxcompose.composer import Composer

TESIS = "/app/tesis_v4.docx"
MANUAL = "/app/Manual_de_Usuario_Referenciado.docx"
OUT = "/app/Tesis_con_Anexo_Manual.docx"

# ---------- 1. Preparar el manual para que no se mezcle con los índices de la tesis ----------
m = Document(MANUAL)
for viejo, nuevo, tam in (("Heading 1", "Manual Titulo 1", 14), ("Heading 2", "Manual Titulo 2", 12)):
    s = m.styles.add_style(nuevo, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = m.styles["Normal"]
    s.font.name = "Arial"; s.font.size = Pt(tam); s.font.bold = True
    s.font.color.rgb = RGBColor.from_string("1F3A5F")
    s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.keep_with_next = True
    for p in m.paragraphs:
        if p.style.name == viejo:
            p.style = s
for viejo, nuevo in (("Cap_Figura", "Manual Figura"), ("Cap_Tabla", "Manual Tabla")):
    s = m.styles[viejo]
    s.name = nuevo
    s.element.set(qn("w:styleId"), nuevo.replace(" ", ""))
    for p in m.element.body.iter(qn("w:pStyle")):
        if p.get(qn("w:val")) == viejo:
            p.set(qn("w:val"), nuevo.replace(" ", ""))
for it in m.element.body.iter(qn("w:instrText")):
    t = it.text or ""
    t = t.replace('\\o "1-2"', '\\t "Manual Titulo 1,1,Manual Titulo 2,2"')
    t = t.replace('"Cap_Figura,1"', '"Manual Figura,1"').replace('"Cap_Tabla,1"', '"Manual Tabla,1"')
    t = re.sub(r"_Toc(\d+)", r"_TocM\1", t)
    it.text = t
for el in m.element.body.iter():
    for atr in (qn("w:name"), qn("w:anchor")):
        v = el.get(atr)
        if v and v.startswith("_Toc"):
            el.set(atr, "_TocM" + v[4:])
# Sin encabezado/pie propios: se usan los de la tesis
sect = m.element.body.find(qn("w:sectPr"))
for ref in sect.findall(qn("w:headerReference")) + sect.findall(qn("w:footerReference")):
    sect.remove(ref)
tp = sect.find(qn("w:titlePg"))
if tp is not None:
    sect.remove(tp)
m.save("/tmp/manual_prep.docx")

# ---------- 2. Insertar en ANEXOS ----------
t = Document(TESIS)
body = t.element.body
kids = list(body)
i_anexos = next(i for i, e in enumerate(kids)
                if e.tag == qn("w:p") and "".join(x.text or "" for x in e.iter(qn("w:t"))).strip() == "ANEXOS")
ancla = kids[i_anexos + 2]          # párrafo vacío previo al salto de página de las notas
assert not "".join(ancla.itertext()).strip()

# Subtítulo del anexo (estilo Título 2 de la tesis) seguido de salto de página
sub = t.add_paragraph("ANEXO A. MANUAL DE USUARIO DE LA APLICACIÓN “TURNOS DE POZO”", style="Heading 2")._p
p_intro = t.add_paragraph(
    "A continuación se presenta el manual de usuario de la aplicación móvil “Turnos de Pozo”, elaborado para "
    "guiar a socios, contadores y usuarios en general en el uso de cada una de sus funciones.")._p
p_intro.get_or_add_pPr()
# Sin numeración automática del Título 2 (evita "4.2 ANEXO A")
pPr = sub.get_or_add_pPr()
numPr = OxmlElement("w:numPr")
for tag, val in (("w:ilvl", "0"), ("w:numId", "0")):
    e = OxmlElement(tag); e.set(qn("w:val"), val); numPr.append(e)
pPr.find(qn("w:pStyle")).addnext(numPr)

# Números de página guardados en los índices internos del manual (se ajustan al nuevo lugar)
OFFSET = 61
mb = Document("/tmp/manual_prep.docx")
for p in mb.paragraphs:
    if p.style.name.startswith("toc") or p.style.name == "Indice_Figuras_Tablas":
        ts = [x for x in p._p.iter(qn("w:t")) if x.text and x.text.strip()]
        if ts and ts[-1].text.strip().isdigit():
            ts[-1].text = str(int(ts[-1].text.strip()) + OFFSET)
mb.save("/tmp/manual_prep.docx")

st = t.settings.element
uf = st.find(qn("w:updateFields"))
if uf is None:
    uf = OxmlElement("w:updateFields")
    sig = next((st.find(qn(x)) for x in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat",
                                          "w:docVars", "w:rsids") if st.find(qn(x)) is not None), None)
    sig.addprevious(uf) if sig is not None else st.append(uf)
uf.set(qn("w:val"), "true")
salto = OxmlElement("w:p"); r = OxmlElement("w:r"); br = OxmlElement("w:br"); br.set(qn("w:type"), "page")
r.append(br); salto.append(r)
ancla.addnext(sub); sub.addnext(p_intro); p_intro.addnext(salto)

idx = list(body).index(salto) + 1
comp = Composer(t)
comp.insert(idx, Document("/tmp/manual_prep.docx"))
comp.save(OUT)
print("ok")
