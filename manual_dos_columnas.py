"""Manual con pasos en dos columnas: instrucciones a la izquierda, captura a la derecha.
Uso: python manual_dos_columnas.py entrada.docx salida.docx"""
import sys, copy, re
from docx import Document
from docx.shared import Emu, Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC, DST = sys.argv[1], sys.argv[2]
ALTO = Emu(int(9.0 * 360000))      # alto de la captura
COL_TXT, COL_IMG = Cm(9.6), Cm(6.4)
AZUL = RGBColor(0x1F, 0x3A, 0x5F)

d = Document(SRC)
body = d.element.body


def sin_bordes(tbl):
    tblPr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'nil')
        b.append(e)
    tblPr.append(b)


def es_figura(el):
    return el is not None and el.tag == qn('w:tbl') and 'graphic' in el.xml and len(el.findall(qn('w:tr'))) == 2


for p in list(d.paragraphs):
    m = re.match(r'Paso (\d+)\.', p.text.strip())
    fig = p._p.getnext()
    if not m or not es_figura(fig):
        continue
    filas = fig.findall(qn('w:tr'))
    p_img = filas[0].find('.//' + qn('w:p'))
    p_cap = filas[1].find('.//' + qn('w:p'))

    # Ajustar tamaño de la captura
    ext = p_img.find('.//' + qn('wp:extent'))
    cx, cy = int(ext.get('cx')), int(ext.get('cy'))
    nuevo_cx = int(cx * ALTO / cy)
    for e in p_img.iter(qn('wp:extent'), qn('a:ext')):
        e.set('cx', str(nuevo_cx)); e.set('cy', str(int(ALTO)))

    # Tabla de 2 columnas
    t = d.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    sin_bordes(t._tbl)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))
    c_txt, c_img = t.rows[0].cells
    c_txt.width, c_img.width = COL_TXT, COL_IMG
    c_txt.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    c_img.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    # Columna izquierda: número grande + instrucciones
    num = c_txt.paragraphs[0]
    r = num.add_run(m.group(1))
    r.bold = True; r.font.size = Pt(32); r.font.color.rgb = AZUL
    num.paragraph_format.space_after = Pt(0)
    p_txt = copy.deepcopy(p._p)
    c_txt._tc.append(p_txt)
    pPr = p_txt.find(qn('w:pPr'))
    kn = pPr.find(qn('w:keepNext')) if pPr is not None else None
    if kn is not None:
        pPr.remove(kn)
    primer = p_txt.find(qn('w:r'))
    primer.find(qn('w:t')).text = 'Paso ' + m.group(1) + '. '

    # Columna derecha: captura + pie de figura
    c_img._tc.remove(c_img.paragraphs[0]._p)
    c_img._tc.append(p_img)
    c_img._tc.append(p_cap)

    # Reemplazar el párrafo y la figura originales por la nueva tabla
    p._p.addprevious(t._tbl)
    body.remove(fig)
    p._p.getparent().remove(p._p)

# Párrafos vacíos tras cada tabla de pasos: separación mínima
for tbl in body.iter(qn('w:tbl')):
    sig = tbl.getnext()
    if sig is not None and sig.tag == qn('w:p') and 'graphic' in tbl.xml and not ''.join(sig.itertext()).strip() and 'w:br' not in sig.xml:
        pf = sig.get_or_add_pPr()
        sp = pf.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing'); pf.append(sp)
        sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '120')
        sp.set(qn('w:line'), '240'); sp.set(qn('w:lineRule'), 'auto')

# Quitar vacíos duplicados consecutivos
for p in list(d.paragraphs):
    prev = p._p.getprevious()
    if (not p.text.strip() and prev is not None and prev.tag == qn('w:p')
            and not ''.join(prev.itertext()).strip()
            and 'w:br' not in p._p.xml + prev.xml and 'graphic' not in p._p.xml + prev.xml
            and 'sectPr' not in p._p.xml + prev.xml):
        pp = prev.getprevious()
        if pp is not None and pp.tag == qn('w:tbl'):
            p._p.getparent().remove(p._p)

d.save(DST)
print('OK', DST)
